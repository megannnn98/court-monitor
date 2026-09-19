"""The `/export` spreadsheet: one row per person and article.

Written the way `web/exports.py` writes the candidates file: scraped names, titles and
URLs are stored as text, so a value starting with `=` never becomes a formula, and only
an http(s) link is made clickable.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from telegram_bot.models import NewsArticleReference, PeopleFromNewsResult

SHEET_TITLE = "Люди из новостей"
HEADERS = (
    "№",
    "Человек",
    "Публикаций за период",
    "Последняя публикация",
    "Дата статьи",
    "Источник",
    "Заголовок",
    "Ссылка",
)
COLUMN_WIDTHS = (5, 34, 12, 20, 12, 22, 70, 50)
DATE_FORMAT = "DD.MM.YYYY"
DATETIME_FORMAT = "DD.MM.YYYY HH:MM"


def people_filename(result: PeopleFromNewsResult) -> str:
    return f"people-{result.date_from.isoformat()}-{result.date_to.isoformat()}.xlsx"


def people_xlsx(result: PeopleFromNewsResult, timezone: ZoneInfo) -> bytes:
    """One row per (person, article); a person without articles still gets a row."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = SHEET_TITLE
    sheet.append(list(HEADERS))
    for column, width in enumerate(COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.cell(row=1, column=column).font = Font(bold=True)
    sheet.freeze_panes = "A2"

    row = 1
    for position, person in enumerate(result.people, start=1):
        latest = _local(person.latest_published_at, timezone)
        # A person is always one row at least, even with no article in the period.
        articles: list[NewsArticleReference | None] = list(person.articles) or [None]
        for article in articles:
            row += 1
            published = None if article is None else _local(article.published_at, timezone)
            sheet.append(
                [
                    position,
                    person.canonical_name,
                    person.article_count,
                    latest,
                    published,
                    None if article is None else article.source_name,
                    None if article is None else article.title,
                    None if article is None else article.url,
                ]
            )
            # Scraped text stays text: a leading "=" must not become a formula.
            for column in (2, 6, 7, 8):
                sheet.cell(row=row, column=column).data_type = "s"
            sheet.cell(row=row, column=4).number_format = DATETIME_FORMAT
            sheet.cell(row=row, column=5).number_format = DATE_FORMAT
            sheet.cell(row=row, column=7).alignment = Alignment(vertical="top", wrap_text=True)
            url = None if article is None else article.url
            if url is not None and url.startswith(("http://", "https://")):
                sheet.cell(row=row, column=8).hyperlink = url
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _local(value: datetime, timezone: ZoneInfo) -> datetime:
    # Excel has no timezone: the moment is converted, then written naive.
    return value.astimezone(timezone).replace(tzinfo=None)
