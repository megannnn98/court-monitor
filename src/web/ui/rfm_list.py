"""«Перечень»: the list itself, and the entries added in a period.

Ирина looks at the перечень for one thing: who was added lately, because a fresh entry
means a fresh case. The state publishes who is in the list and says nothing about when,
so this page is where the days of inclusion ОВД-Инфо publish become usable — and the
filter is «who was added between these dates».

What a row here is: **an entry of the перечень**. Not a person, and not anybody our news
are about. A name on this page has been matched to nobody; it is the list's own record.
The date beside it says when that record appeared, and it is the strongest thing we can
say about it. The wording throughout is about the entry, and Ирина is told so on the
page — this is the page where that is most likely to be read as a statement about people.

The export carries the same rows under the same filter, because she has no access to the
site and works from files.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, date, datetime
from html import escape
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Depends, Query, Response
from fastapi.responses import HTMLResponse
from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from db.orm_models import (
    RosfinmonitoringEntryRecord,
)
from entities.rf_check import FULL
from rosfinmonitoring.inclusion_dates import ATTRIBUTION
from rosfinmonitoring.snapshot_lookup import (
    RosfinmonitoringSnapshotSummary,
    SqlAlchemyRosfinmonitoringSnapshotLookup,
)
from web.dependencies import get_db, session_factory_for
from web.ui.layout import _page

router = APIRouter()

PAGE_SIZE = 100
# The table's columns, for the empty row that has to span all of them.
COLUMNS = 5

# The periods offered as buttons. Days, not months: the list is added to daily, and the
# period an operator asks for is «who came in since Tuesday».
PERIOD_DAYS = ((7, "7 дней"), (30, "30 дней"), (90, "90 дней"), (365, "год"))

_INTRO = (
    "Список террористов и экстремистов Росфинмониторинга — по данным ОВД-Инфо, которые "
    "ведут его историю. Дата в строке — это день, когда запись появилась в перечне: она "
    "описывает запись, а не человека, и не делает совпадение имени подтверждением "
    "личности. Свежая запись — повод посмотреть, не появилось ли по этому имени дело."
)


@dataclass(frozen=True)
class ListFilters:
    """A period of inclusion: days, or two dates. Empty means no period chosen."""

    days: int = 0
    date_from: date | None = None
    date_to: date | None = None

    @property
    def custom(self) -> bool:
        return self.date_from is not None or self.date_to is not None


@dataclass
class ListRow:
    """One entry of the list: who, born when, and when the record appeared."""

    entry_id: int
    full_name: str
    birth_date: date | None
    birth_place: str | None
    inclusion_date: date | None
    # Our own match, if this entry was matched to a person in a case. Said as a link
    # because it is a claim we could not establish from the name alone, and the operator
    # is the one who decides what to do with it.
    group_key: str | None = None
    group_name: str | None = None
    # What the comparison established about that match: `full` (the name with the
    # patronymic) or `name`, which may be a namesake. A namesake is a different person
    # and must not read like the one we did identify.
    match_level: str | None = None


def _latest(db: Session) -> RosfinmonitoringSnapshotSummary | None:
    """The published list we are working against, never the operator's own Airtable one.

    The Airtable snapshot is excluded inside the lookup: a list an operator curates is
    not the list the state publishes, and a page headed «Перечень РФМ» must not show it.
    """
    return SqlAlchemyRosfinmonitoringSnapshotLookup(
        session_factory_for(db)
    ).latest_imported_snapshot()


def _day(moment: datetime | None) -> str:
    return moment.astimezone().strftime("%d.%m.%Y") if moment else "—"


def _period(chosen: ListFilters, today: date) -> tuple[date | None, date | None]:
    """The dates a period means: `(from, to)`, either of which may be open."""
    if chosen.date_from or chosen.date_to:
        return chosen.date_from, chosen.date_to
    if chosen.days:
        return today - date.resolution * chosen.days, today
    return None, None


def _entries(
    db: Session,
    snapshot_id: int,
    chosen: ListFilters,
    today: date,
    *,
    limit: int | None = PAGE_SIZE,
) -> list[ListRow]:
    """The entries of one snapshot under a period, most recently added first.

    The period and the limit are read inside the query, not in Python: this list is 23 023
    rows and the page shows a hundred of them. `limit=None` is for the export, which
    holds every row of the period — that is the file the operator works from.
    """
    start, end = _period(chosen, today)
    query = select(RosfinmonitoringEntryRecord).where(
        RosfinmonitoringEntryRecord.snapshot_id == snapshot_id
    )
    if start is not None:
        query = query.where(RosfinmonitoringEntryRecord.inclusion_date >= _moment(start))
    if end is not None:
        # The whole last day: `inclusion_date` is a day, and `<= 00:00` would drop it.
        query = query.where(
            RosfinmonitoringEntryRecord.inclusion_date < _moment(end + date.resolution)
        )
    if start is None and end is None:
        # With no period chosen the list is ordered by the day we know, so the entries
        # that have one are the ones worth seeing.
        query = query.where(RosfinmonitoringEntryRecord.inclusion_date.isnot(None))
    query = query.order_by(
        RosfinmonitoringEntryRecord.inclusion_date.desc().nulls_last(),
        RosfinmonitoringEntryRecord.full_name,
    )
    if limit is not None:
        query = query.limit(limit)
    return [
        ListRow(
            entry_id=entry.id,
            full_name=entry.full_name,
            birth_date=entry.birth_date.date() if entry.birth_date else None,
            birth_place=entry.birth_place,
            inclusion_date=entry.inclusion_date.date() if entry.inclusion_date else None,
        )
        for entry in db.scalars(query)
    ]


def _moment(day: date) -> datetime:
    return datetime(day.year, day.month, day.day, tzinfo=UTC)


def _with_our_matches(db: Session, rows: list[ListRow]) -> None:
    """Fill in, for each entry, the person our comparison matched it to — if any.

    A link, not a verdict: the comparison matched a name, and the day on the row is the
    entry's whatever that match was. How well it matched is kept, because a namesake is
    a different person and has to read as one.

    Only the entities the page actually shows are read: an earlier version loaded every
    group in the database to look up a handful of links, which on a real database is
    tens of thousands of rows read to fill in two columns of a hundred rows.
    """
    from db.orm_models import EntityGroupRecord, EntityGroupRfMatchRecord

    if not rows:
        return
    by_entry = {row.entry_id: row for row in rows}
    for entry_id, group_id, level, key, name in db.execute(
        select(
            EntityGroupRfMatchRecord.entry_id,
            EntityGroupRfMatchRecord.group_id,
            EntityGroupRfMatchRecord.level,
            EntityGroupRecord.key,
            EntityGroupRecord.name,
        )
        .select_from(EntityGroupRfMatchRecord)
        .join(EntityGroupRecord, EntityGroupRecord.id == EntityGroupRfMatchRecord.group_id)
        .where(EntityGroupRfMatchRecord.entry_id.in_(list(by_entry)))
        # `level = 'full'` first, as on «Результате» — and not `level DESC`, which sorts
        # the two words alphabetically and puts «name» ahead of «full», so where one
        # entry matched two people the namesake would be the one shown.
        .order_by((EntityGroupRfMatchRecord.level == FULL).desc())
    ).all():
        row = by_entry[entry_id]
        if row.group_key is not None:
            continue
        row.group_key, row.group_name, row.match_level = key, name, level


def _match_text(row: ListRow) -> str:
    """What the comparison established, in the words used everywhere else.

    «в перечне» for a name matched with its patronymic, «возможно тёзка» for one matched
    on the name and surname alone. The link looks the same in both, so the distinction
    has to be in the words or the operator cannot see it.
    """
    if row.group_key is None:
        return ""
    return "в перечне" if row.match_level == FULL else "возможно тёзка"


def _count(db: Session, snapshot_id: int, chosen: ListFilters, today: date) -> tuple[int, int, int]:
    """(in the period, with a date, in all): the three numbers the page opens with."""
    start, end = _period(chosen, today)
    dated = db.scalar(
        select(func.count())
        .select_from(RosfinmonitoringEntryRecord)
        .where(
            RosfinmonitoringEntryRecord.snapshot_id == snapshot_id,
            RosfinmonitoringEntryRecord.inclusion_date.isnot(None),
        )
    )
    query = (
        select(func.count())
        .select_from(RosfinmonitoringEntryRecord)
        .where(
            RosfinmonitoringEntryRecord.snapshot_id == snapshot_id,
            RosfinmonitoringEntryRecord.inclusion_date.isnot(None),
        )
    )
    if start is not None:
        query = query.where(RosfinmonitoringEntryRecord.inclusion_date >= _moment(start))
    if end is not None:
        query = query.where(
            RosfinmonitoringEntryRecord.inclusion_date < _moment(end + date.resolution)
        )
    in_period = db.scalar(query) or 0
    total = (
        db.scalar(
            select(func.count())
            .select_from(RosfinmonitoringEntryRecord)
            .where(RosfinmonitoringEntryRecord.snapshot_id == snapshot_id)
        )
        or 0
    )
    return int(in_period), int(dated or 0), int(total)


def _filters_html(chosen: ListFilters, today: date) -> str:
    """The period buttons and the two date fields, with the same names the export reads."""
    periods = " ".join(
        f'<button type="submit" name="days" value="{value}" '
        f'class="chip{" active" if value == chosen.days and not chosen.custom else ""}" '
        "onclick=\"this.form.date_from.value='';this.form.date_to.value=''\">"
        f"{label}</button>"
        for value, label in PERIOD_DAYS
    )
    return f"""<form method="get" action="/ui/rfm" class="toolbar" id="rfm-filters">
  <input type="hidden" name="days" value="{chosen.days}">
  <span class="chips">{periods}</span>
  <span class="chips">
    <label class="field">Включены с
      <input type="date" name="date_from" value="{chosen.date_from or ""}"></label>
    <label class="field">по
      <input type="date" name="date_to" value="{chosen.date_to or ""}"></label>
    <button type="submit">Показать</button>
  </span>
  <div class="filter-row">
    <button type="submit" class="secondary" formaction="/ui/rfm/export.xlsx">Скачать Excel</button>
  </div>
</form>"""


def _rows_html(rows: list[ListRow]) -> str:
    if not rows:
        return f'<tr><td colspan="{COLUMNS}" class="muted">Записей нет.</td></tr>'
    cells = []
    for row in rows:
        matched = (
            f'<a href="/ui/investigations/{quote(row.group_key)}">{escape(row.group_name or "")}</a>'
            f'<div class="muted">{escape(_match_text(row))}</div>'
            if row.group_key
            else '<span class="muted">—</span>'
        )
        cells.append(
            f"<tr><td>{escape(row.full_name)}</td>"
            f"<td>{f'{row.birth_date:%d.%m.%Y}' if row.birth_date else '—'}</td>"
            f"<td>{escape(row.birth_place or '—')}</td>"
            f"<td>{f'{row.inclusion_date:%d.%m.%Y}' if row.inclusion_date else '—'}</td>"
            f"<td>{matched}</td></tr>"
        )
    return "".join(cells)


@router.get("/ui/rfm", response_class=HTMLResponse)
def ui_rfm(
    days: int = Query(default=0, ge=0, le=3650),
    date_from: str = Query(default="", max_length=10),
    date_to: str = Query(default="", max_length=10),
    db: Session = Depends(get_db),  # noqa: B008
) -> HTMLResponse:
    """The list, and the entries added in a period."""
    today = datetime.now(UTC).date()
    chosen = ListFilters(days, _parse_day(date_from), _parse_day(date_to))
    latest = _latest(db)
    if latest is None:
        return _page(
            "Перечень РФМ",
            '<p class="empty">Перечень ещё не загружен.</p>',
            active="rfm",
            instruction=_INTRO,
            db=db,
        )
    rows = _entries(db, latest.snapshot_id, chosen, today)
    _with_our_matches(db, rows)
    in_period, dated, total = _count(db, latest.snapshot_id, chosen, today)
    period_note = f"За период: {in_period}." if in_period else "За выбранный период записей нет."
    body = f"""{_filters_html(chosen, today)}
<p class="muted">Снимок перечня от {_day(latest.snapshot_date)}: {total} записей, из них
с датой включения {dated}. {period_note} Показаны первые {len(rows)}; файл содержит все.
Дата — свойство записи перечня, а не человека: совпадение имени и даты рождения не
подтверждает, что в новости речь о том же человеке. {escape(ATTRIBUTION)}.</p>
<table><thead><tr><th>ФИО в перечне</th><th>Дата рождения</th><th>Место рождения</th>
<th>Запись включена</th><th>Наше совпадение</th></tr></thead>
<tbody>{_rows_html(rows)}</tbody></table>"""
    return _page("Перечень РФМ", body, active="rfm", instruction=_INTRO, db=db)


def _parse_day(text: str) -> date | None:
    try:
        return date.fromisoformat(text) if text else None
    except ValueError:
        return None


def _export_name(chosen: ListFilters, today: date) -> str:
    """`perechen_<from>_<to>.xlsx`: the period the file holds, so a file in a folder says
    what it is without opening it."""
    start, end = _period(chosen, today)
    if start is None and end is None:
        return "perechen_vse.xlsx"
    first = start.isoformat() if start else "do"
    last = end.isoformat() if end else today.isoformat()
    return f"perechen_{first}_{last}.xlsx"


def rfm_xlsx(rows: list[ListRow]) -> bytes:
    """The list as an Excel file, the way «Результат» writes its own.

    A CSV would be a file Russian Excel opens in one column, and the operator works from
    files rather than from the site — so the file is the deliverable, and a deliverable
    that opens wrongly is not one. Dates go in as dates, not as text, so a period can be
    sorted and counted in the file itself.
    """
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Перечень"
    sheet.append(
        [
            "ФИО в перечне",
            "Дата рождения",
            "Место рождения",
            "Запись включена",
            "Наше совпадение",
            "Что это значит",
        ]
    )
    for row in rows:
        sheet.append(
            [
                row.full_name,
                row.birth_date,
                row.birth_place or None,
                row.inclusion_date,
                row.group_name or None,
                _match_text(row) or None,
            ]
        )
    for line in range(2, len(rows) + 2):
        # Names and places come from a scraped page: never let «=» become a formula.
        for column in (1, 3, 5, 6):
            cell = sheet.cell(row=line, column=column)
            if cell.value is not None:
                cell.data_type = "s"
        for column in (2, 4):
            sheet.cell(row=line, column=column).number_format = "DD.MM.YYYY"
    for letter, width in (("A", 42), ("B", 18), ("C", 40), ("D", 20), ("E", 30), ("F", 20)):
        sheet.column_dimensions[letter].width = width
    # The attribution and the rule travel with the rows: the file leaves the site.
    note = workbook.create_sheet("Источник дат")
    note["A1"] = "Дата включения в перечень"
    note["A2"] = ATTRIBUTION
    note["A3"] = (
        "Дата описывает запись перечня, а не человека: совпадение имени и даты рождения "
        "не подтверждает, что в новости речь о том же человеке."
    )
    note["A4"] = (
        "«Возможно тёзка» — совпали имя и фамилия, отчества нет с одной стороны: это может "
        "быть другой человек, и его запись о дате включения ничего о нём не говорит."
    )
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@router.get("/ui/rfm/export.xlsx")
def ui_rfm_export(
    days: int = Query(default=0, ge=0, le=3650),
    date_from: str = Query(default="", max_length=10),
    date_to: str = Query(default="", max_length=10),
    db: Session = Depends(get_db),  # noqa: B008
) -> Response:
    """The same rows the page shows, under the same period, as an Excel file.

    Excel and the page must not disagree: the filter is read from the same two names, so
    a file downloaded from a filtered page holds exactly that period.
    """
    today = datetime.now(UTC).date()
    chosen = ListFilters(days, _parse_day(date_from), _parse_day(date_to))
    latest = _latest(db)
    rows = _entries(db, latest.snapshot_id, chosen, today, limit=None) if latest is not None else []
    _with_our_matches(db, rows)
    return Response(
        content=rfm_xlsx(rows),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{_export_name(chosen, today)}"'},
    )
