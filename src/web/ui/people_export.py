"""«Выгрузить всех людей»: the whole base of the people as one Excel file.

The result and the people page export what their filters show. This is the other thing: the
operator keeps her own table and wants everybody the system has found, with no filter between
her and the rows — the people with a name on one sheet, and the cases that name nobody
(«15-летний житель Канаша») on another, so that nothing she can see on the site is missing
from the file. She filters in Excel.

Both sheets say as little as the data does. A person nobody has judged has no verdict and no
role, and the cell is empty rather than «не ясно»: an empty cell is «не определено», and
«не ясно» is an answer the model gave.
"""

from __future__ import annotations

from datetime import UTC, datetime
from typing import Any

from fastapi import APIRouter, Depends, Response
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from db.orm_models import (
    EntityGroupChargeRecord,
    EntityGroupNewsRecord,
    EntityGroupPoliticsRecord,
    EntityGroupRecord,
    EntityGroupRoleRecord,
)
from entities.done_marks import done_marks, is_done
from entities.known_base import KnownBase
from entities.news import KIND_LABELS as NEWS_LABELS
from entities.politics import VERDICT_LABELS
from entities.rf_entry import rf_word, strongest_entries
from entities.unnamed import EVENT_LABELS
from entities.unnamed_cases import Case, all_cases
from rosfinmonitoring.inclusion_dates import ATTRIBUTION as INCLUSION_ATTRIBUTION
from web.dependencies import get_db
from web.exports import XLSX, excel_day, is_web_link, workbook_bytes, write_notes, write_sheet
from web.ui.entities import article_order, display_name, regions_text, role_label
from web.ui.political import known_answer_text

router = APIRouter()

# A column: its header and its width.
_PEOPLE_COLUMNS = (
    ("№", 5),
    ("Фамилия Имя", 28),
    ("Роль в деле", 26),
    ("Дело", 14),
    ("Статьи УК", 18),
    ("Регион", 22),
    ("Упоминаний", 11),
    ("Публикаций", 11),
    ("Последняя новость", 16),
    ("Свежая новость", 20),
    ("Перечень РФМ", 50),
    ("В базе Airtable", 36),
    ("Обработано", 12),
)
_SOURCE = "Источник"
_UNNAMED_COLUMNS = (
    ("№", 5),
    ("Кто", 34),
    ("Политическая статья", 12),
    ("Статьи УК", 18),
    ("Место", 22),
    ("Предложений", 11),
    ("Первая новость", 16),
    ("Последняя новость", 16),
    ("Событие", 18),
    ("Что сказано", 60),
    (_SOURCE, 50),
    ("Обработано", 12),
)


def notes_text(stamp: str) -> list[str]:
    return [
        f"Выгрузка всей базы людей на {stamp}: без фильтров.",
        (
            "«Люди» — все, кого система нашла по имени. «Без имени» — случаи, где публикация "
            "не называет человека; опознанные в этот лист не входят: они в «Людях»."
        ),
        (
            "Пустая «Роль в деле» или «Дело» — шаг ещё не определял; «не ясно» — модель не "
            "смогла решить."
        ),
        (
            "Дата в «Перечне РФМ» — свойство записи перечня, а не человека: она не делает "
            "совпадение подтверждением личности. " + INCLUSION_ATTRIBUTION
        ),
        (
            "«В базе Airtable» — сверка только по имени (даты рождения и региона там нет): "
            "«вероятно» и «тёзки» — не уверенность."
        ),
    ]


def people_rows(db: Session) -> list[list[Any]]:
    """Every person with a name, the most mentioned first."""
    groups = list(
        db.scalars(
            select(EntityGroupRecord).order_by(
                EntityGroupRecord.mention_count.desc(), EntityGroupRecord.key
            )
        )
    )
    ids = [group.id for group in groups]
    roles = {
        group_id: (role, kind)
        for group_id, role, kind in db.execute(
            select(
                EntityGroupRoleRecord.group_id,
                EntityGroupRoleRecord.role,
                EntityGroupRoleRecord.kind,
            )
        )
    }
    verdicts = {
        group_id: verdict
        for group_id, verdict in db.execute(
            select(EntityGroupPoliticsRecord.group_id, EntityGroupPoliticsRecord.verdict)
        )
    }
    news = {
        group_id: kind
        for group_id, kind in db.execute(
            select(EntityGroupNewsRecord.group_id, EntityGroupNewsRecord.kind)
        )
    }
    articles: dict[int, set[str]] = {}
    for group_id, article in db.execute(
        select(EntityGroupChargeRecord.group_id, EntityGroupChargeRecord.article).distinct()
    ):
        articles.setdefault(group_id, set()).add(article)
    entries = strongest_entries(db, ids)
    base = KnownBase.from_session(db)
    marks = done_marks(db)
    rows: list[list[Any]] = []
    for position, group in enumerate(groups, start=1):
        role = roles.get(group.id)
        entry = entries.get(group.id)
        rows.append(
            [
                position,
                display_name(group.name),
                role_label(*role) if role else None,
                VERDICT_LABELS.get(verdicts.get(group.id, "")),
                ", ".join(sorted(articles.get(group.id, ()), key=article_order)) or None,
                regions_text(group.regions) or None,
                group.mention_count,
                group.article_count,
                excel_day(group.last_published_at),
                NEWS_LABELS.get(news[group.id], news[group.id]) if group.id in news else None,
                rf_word(entry.level, entry.text, entry.inclusion_date, entry.inclusion_source)
                if entry
                else None,
                known_answer_text(base.match(group.name), loaded=bool(len(base))),
                "да" if is_done(marks, group.key, group.last_published_at) else None,
            ]
        )
    return rows


def _case_source(case: Case) -> str:
    latest = case.latest
    return f"{latest.source}: {latest.title[:120]}" if latest.title else latest.source


def unnamed_rows(db: Session) -> tuple[list[list[Any]], list[str]]:
    """Every unnamed person not yet identified, the latest news first; and the addresses
    of their latest sources, to make each a link."""
    marks = done_marks(db)
    oldest = datetime.min.replace(tzinfo=UTC)
    cases = sorted(all_cases(db), key=lambda case: case.last_published_at or oldest, reverse=True)
    rows = [
        [
            position,
            case.name,
            "да" if case.political else "нет",
            ", ".join(case.articles) or None,
            case.place or None,
            len(case.sentences),
            excel_day(case.first_published_at),
            excel_day(case.last_published_at),
            EVENT_LABELS.get(case.latest.event_type, case.latest.event_type),
            case.latest.explanation,
            _case_source(case),
            "да" if is_done(marks, case.key, case.last_published_at) else None,
        ]
        for position, case in enumerate(cases, start=1)
    ]
    return rows, [case.latest.url for case in cases]


def _table(sheet: Worksheet, columns: tuple[tuple[str, int], ...], rows: list[list[Any]]) -> None:
    write_sheet(
        sheet,
        [header for header, _ in columns],
        rows,
        widths=[width for _, width in columns],
        filterable=True,
    )


def people_xlsx(db: Session, *, today: datetime | None = None) -> bytes:
    workbook = Workbook()
    people = workbook.active
    assert people is not None
    people.title = "Люди"
    _table(people, _PEOPLE_COLUMNS, people_rows(db))
    unnamed = workbook.create_sheet("Без имени")
    rows, links = unnamed_rows(db)
    _table(unnamed, _UNNAMED_COLUMNS, rows)
    source = [header for header, _ in _UNNAMED_COLUMNS].index(_SOURCE) + 1
    for line, url in enumerate(links, start=2):
        if is_web_link(url):
            unnamed.cell(row=line, column=source).hyperlink = url
    stamp = (today or datetime.now(UTC)).astimezone().strftime("%d.%m.%Y %H:%M")
    write_notes(workbook, "Пояснения", notes_text(stamp))
    workbook["Пояснения"].column_dimensions["A"].width = 140
    return workbook_bytes(workbook)


# Not under `/ui/entities/`: there `/ui/entities/{key}` is a person's page, and «export.xlsx»
# would be taken for a key.
@router.get("/ui/people/export.xlsx")
def ui_people_export(db: Session = Depends(get_db)) -> Response:  # noqa: B008
    """The whole base of the people, whatever the page's filters say."""
    name = f"people_{datetime.now(UTC).date().isoformat()}.xlsx"
    return Response(
        people_xlsx(db),
        media_type=XLSX,
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )
