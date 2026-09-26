"""«Результат»: the politically persecuted — what the whole pipeline is for.

A figurant of a criminal case (step 4) whose case is political persecution (step 5), on
the Rosfinmonitoring list or not: the list does not make a case known, it confirms who a
person is (the operator's word) — its birth date and place are shown beside the name. A
name the list carries without a patronymic may be a namesake: marked, and can be hidden.
The period filter answers «new or old case» by the date of the latest publication; the
same rows go to Excel.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import UTC, date, datetime, time, timedelta
from html import escape
from io import BytesIO
from typing import Any
from urllib.parse import parse_qs, quote, unquote, urlencode

from fastapi import APIRouter, Depends, Query, Request
from fastapi.responses import HTMLResponse, Response
from openpyxl import Workbook
from sqlalchemy import exists, func, select, text
from sqlalchemy.orm import Session

from db.orm_models import (
    EntityGroupNewsRecord,
    EntityGroupPoliticsRecord,
    EntityGroupRecord,
    EntityGroupRfMatchRecord,
)
from entities.news import KIND_LABELS, NEW_CASE, ONGOING, OTHER, SENTENCE, UNKNOWN
from entities.politics import MEMORIAL_CATEGORIES, POLITICAL
from entities.rf_check import FULL
from persecution.classifier import POLITICAL_ARTICLES
from web.dependencies import get_db
from web.ui.entities import _articles_by_group, _date, display_name
from web.ui.funnel import funnel, funnel_line
from web.ui.layout import _page, pager

router = APIRouter()

PAGE_SIZE = 100
LINKS = 3
PERIODS = {0: "За всё время", 1: "Месяц", 3: "3 месяца", 6: "Полгода", 12: "Год"}
# The last filters, so that a reload or the menu's link keeps the period.
FILTERS_COOKIE = "political_filters"
FILTER_NAMES = ("months", "date_from", "date_to", "hide_maybe_listed", "news")
# What the latest news is (`entities.news`): the operator's new cases and sentences first.
NEWS_FILTERS = {
    "all": "Любая свежая новость",
    NEW_CASE: "Новые дела",
    SENTENCE: "Приговоры",
    ONGOING: "Продолжение дела",
    OTHER: "Другое",
    UNKNOWN: "Не определено",
}

_PUBLICATIONS = text(
    """
    SELECT DISTINCT gm.group_id, a.id, a.title, a.published_at, d.canonical_url
    FROM entity_group_mentions gm
    JOIN entity_mentions m ON m.id = gm.mention_id
    JOIN article_extraction_runs r ON r.id = m.extraction_run_id
    JOIN parsed_articles a ON a.id = r.article_id
    JOIN source_documents d ON d.id = a.document_id
    WHERE gm.group_id = ANY(:groups)
    """
)


_RF_ENTRIES = text(
    """
    SELECT m.group_id, m.level, e.full_name, e.birth_date, e.birth_place
    FROM entity_group_rf_matches m JOIN rosfinmonitoring_entries e ON e.id = m.entry_id
    WHERE m.group_id = ANY(:groups)
    ORDER BY m.group_id, m.level = 'full' DESC, e.full_name
    """
)


def rf_text(row: ListRow) -> str:
    """The list's word on the person, for the page and for Excel."""
    if row.rf_level == FULL:
        return f"в перечне: {row.rf_entry}"
    if row.rf_level is not None:
        return f"возможно тёзка: {row.rf_entry}"
    return ""


@dataclass
class ListRow:
    entity: EntityGroupRecord
    politics: EntityGroupPoliticsRecord
    # On the list by the name alone (no patronymic on one side): maybe a namesake.
    maybe_listed: bool
    articles: list[tuple[str, bool]] = field(default_factory=list)
    # The list's entry: «full» (the name with the patronymic) or «name», and who it is.
    rf_level: str | None = None
    rf_entry: str = ""
    # What the latest news is, and why the model said so.
    news_kind: str | None = None
    news_reason: str = ""
    memorial: str | None = None
    first_published: datetime | None = None
    # (title, url, published) of the latest publications.
    links: list[tuple[str, str, datetime | None]] = field(default_factory=list)


@dataclass(frozen=True)
class Filters:
    """The list's filters: a period of the latest news — the last months, or dates — and
    whether to hide those the list may carry without a patronymic."""

    months: int = 0
    date_from: date | None = None
    date_to: date | None = None
    hide_maybe_listed: bool = False
    news: str = "all"

    @property
    def custom(self) -> bool:
        return self.date_from is not None or self.date_to is not None

    def query(self) -> dict[str, str]:
        """As URL parameters, for the pager and the Excel link."""
        return {
            "months": str(self.months),
            "date_from": self.date_from.isoformat() if self.date_from else "",
            "date_to": self.date_to.isoformat() if self.date_to else "",
            "hide_maybe_listed": str(self.hide_maybe_listed).lower(),
            "news": self.news,
        }


def _parse_date(text: str) -> date | None:
    """«25/09/2026» as the form writes it (day/month/year), «25.09.2026», or the
    ISO «2026-09-25» of the links and the cookie."""
    value = text.strip()
    if not value:
        return None
    for layout in ("%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(value, layout).date()  # noqa: DTZ007 - a date, no time
        except ValueError:
            continue
    return None


def _form_date(value: date | None) -> str:
    return f"{value:%d/%m/%Y}" if value else ""


def _date_field(name: str, label: str, value: date | None) -> str:
    """A day/month/year field, whatever the browser's language: a native date field
    shows the browser's own order (month first in an English one)."""
    return (
        f'<label class="dates">{label} <input type="text" name="{name}" '
        f'value="{_form_date(value)}" placeholder="дд/мм/гггг" inputmode="numeric" '
        'pattern="\\d{1,2}/\\d{1,2}/\\d{4}" title="день/месяц/год" size="10" '
        "data-date></label>"
    )


def filters(
    months: int, date_from: str, date_to: str, hide_maybe_listed: bool, news: str = "all"
) -> Filters:
    """Dates, when given, win over the months."""
    start, end = _parse_date(date_from), _parse_date(date_to)
    return Filters(
        months=0 if start or end or months not in PERIODS else months,
        date_from=start,
        date_to=end,
        hide_maybe_listed=hide_maybe_listed,
        news=news if news in NEWS_FILTERS else "all",
    )


def remembered(cookie: str) -> Filters | None:
    """The filters the cookie keeps; None when it keeps nothing usable."""
    # The page's script writes it encoded; the server, plain.
    values = {name: items[-1] for name, items in parse_qs(unquote(cookie)).items()}
    if not values:
        return None
    try:
        months = int(values.get("months") or 0)
    except ValueError:
        months = 0
    return filters(
        months,
        values.get("date_from", "")[:10],
        values.get("date_to", "")[:10],
        values.get("hide_maybe_listed") == "true",
        values.get("news", "all"),
    )


def _rows(
    db: Session, chosen: Filters
) -> tuple[list[tuple[EntityGroupRecord, EntityGroupPoliticsRecord]], int, int, dict[str, int]]:
    """The list's entities, latest news first; how many there are, how many of them the
    list may carry under a name without a patronymic, and how many of each latest news
    (in the period, before the choice of the news)."""
    # By the name alone and never with the patronymic: maybe a namesake.
    maybe_listed = exists().where(
        EntityGroupRfMatchRecord.group_id == EntityGroupRecord.id,
        EntityGroupRfMatchRecord.level != FULL,
    ) & ~exists().where(
        EntityGroupRfMatchRecord.group_id == EntityGroupRecord.id,
        EntityGroupRfMatchRecord.level == FULL,
    )
    query = (
        select(EntityGroupRecord, EntityGroupPoliticsRecord)
        .join(EntityGroupPoliticsRecord, EntityGroupPoliticsRecord.group_id == EntityGroupRecord.id)
        .where(EntityGroupPoliticsRecord.verdict == POLITICAL)
    )
    if chosen.months:
        since = datetime.now(UTC) - timedelta(days=30 * chosen.months)
        query = query.where(EntityGroupRecord.last_published_at >= since)
    if chosen.date_from is not None:
        start = datetime.combine(chosen.date_from, time.min, UTC)
        query = query.where(EntityGroupRecord.last_published_at >= start)
    if chosen.date_to is not None:
        # The whole last day.
        end = datetime.combine(chosen.date_to + timedelta(days=1), time.min, UTC)
        query = query.where(EntityGroupRecord.last_published_at < end)
    maybe = db.scalar(select(func.count()).select_from(query.where(maybe_listed).subquery())) or 0
    if chosen.hide_maybe_listed:
        query = query.where(~maybe_listed)
    in_period = query.subquery()
    news_counts = {
        str(kind): count
        for kind, count in db.execute(
            select(EntityGroupNewsRecord.kind, func.count())
            .join(in_period, in_period.c.id == EntityGroupNewsRecord.group_id)
            .group_by(EntityGroupNewsRecord.kind)
        ).all()
    }
    if chosen.news != "all":
        query = query.where(
            exists().where(
                EntityGroupNewsRecord.group_id == EntityGroupRecord.id,
                EntityGroupNewsRecord.kind == chosen.news,
            )
        )
    rows = db.execute(
        query.order_by(
            EntityGroupRecord.last_published_at.desc().nulls_last(), EntityGroupRecord.key
        )
    ).all()
    return [(entity, politics) for entity, politics in rows], len(rows), maybe, news_counts


def _details(
    db: Session, found: list[tuple[EntityGroupRecord, EntityGroupPoliticsRecord]]
) -> list[ListRow]:
    ids = [entity.id for entity, _ in found]
    rows = {
        entity.id: ListRow(entity=entity, politics=politics, maybe_listed=False)
        for entity, politics in found
    }
    for group_id, kind, reason in db.execute(
        select(
            EntityGroupNewsRecord.group_id, EntityGroupNewsRecord.kind, EntityGroupNewsRecord.reason
        ).where(EntityGroupNewsRecord.group_id.in_(ids))
    ).all():
        rows[group_id].news_kind, rows[group_id].news_reason = kind, reason
    # The strongest entry of each: the query gives those with the patronymic first.
    for group_id, level, full_name, birth_date, birth_place in db.execute(
        _RF_ENTRIES, {"groups": ids}
    ).all():
        row = rows[group_id]
        if row.rf_level is not None:
            continue
        row.rf_level = level
        row.rf_entry = ", ".join(
            part
            for part in (
                full_name,
                f"{birth_date:%d.%m.%Y} г.р." if birth_date else "",
                birth_place or "",
            )
            if part
        )
    for row in rows.values():
        row.maybe_listed = row.rf_level is not None and row.rf_level != FULL
    for group_id, articles in _articles_by_group(db, ids).items():
        rows[group_id].articles = articles
    for group_id, category in db.execute(MEMORIAL_CATEGORIES, {"groups": ids}).all():
        rows[group_id].memorial = category
    publications: dict[int, list[tuple[str, str, datetime | None]]] = {}
    for group_id, _id, title, published_at, url in db.execute(_PUBLICATIONS, {"groups": ids}).all():
        publications.setdefault(group_id, []).append((title, url, published_at))
    oldest = datetime.min.replace(tzinfo=UTC)
    for group_id, found_publications in publications.items():
        dates = [published for _, _, published in found_publications if published is not None]
        rows[group_id].first_published = min(dates) if dates else None
        rows[group_id].links = sorted(
            found_publications, key=lambda item: item[2] or oldest, reverse=True
        )[:LINKS]
    return [rows[entity.id] for entity, _ in found]


def _article_text(articles: list[tuple[str, bool]]) -> str:
    return ", ".join(article for article, _ in articles)


def _regions_text(entity: EntityGroupRecord) -> str:
    return ", ".join(str(region) for region, _ in entity.regions)


def _basis(row: ListRow) -> str:
    method = "статья УК" if row.politics.method == "article" else "модель"
    return f"{method}: {row.politics.reason}"


def _news_mark(row: ListRow) -> str:
    """The latest news as a mark; the new cases and the sentences stand out."""
    if row.news_kind is None:
        return '<span class="muted">—</span>'
    css = {NEW_CASE: "succeeded", SENTENCE: "succeeded", UNKNOWN: "pending"}.get(row.news_kind, "")
    return (
        f'<span class="badge {css}" title="{escape(row.news_reason, quote=True)}">'
        f"{escape(KIND_LABELS.get(row.news_kind, row.news_kind))}</span>"
    )


def _html_row(position: int, row: ListRow) -> str:
    entity = row.entity
    articles = ", ".join(
        f"<b>{escape(article)}</b>" if article in POLITICAL_ARTICLES else escape(article)
        for article, _ in row.articles
    )
    links = "<br>".join(
        f'<a href="{escape(url, quote=True)}">{escape(title[:80])}</a>'
        for title, url, _ in row.links
        if url.startswith(("http://", "https://"))
    )
    listed = (
        ' <span class="badge">в перечне РФМ</span>'
        if row.rf_level == FULL
        else ' <span class="badge pending">возможно в перечне</span>'
        if row.maybe_listed
        else ""
    )
    return (
        f"<tr><td>{position}</td>"
        f'<td><a href="/ui/investigations/{quote(entity.key)}">'
        f"{escape(display_name(entity.name))}</a>{listed}</td>"
        f"<td>{_news_mark(row)}</td>"
        f"<td>{escape(_regions_text(entity))}</td>"
        f'<td class="muted">{escape(row.rf_entry)}</td>'
        f"<td>{articles}</td>"
        f'<td>{escape(_basis(row))}<br><span class="muted">{escape(row.politics.quote[:200])}</span></td>'
        f"<td>{escape(row.memorial or '')}</td>"
        f"<td>{escape(_date(row.first_published))}</td>"
        f"<td>{escape(_date(entity.last_published_at))}</td>"
        f"<td>{links}</td></tr>"
    )


@router.get("/ui/political", response_class=HTMLResponse)
def ui_political(
    request: Request,
    months: int = Query(default=0),
    date_from: str = Query(default="", max_length=10),
    date_to: str = Query(default="", max_length=10),
    hide_maybe_listed: bool = Query(default=False),
    news: str = Query(default="all", max_length=16),
    page: int = Query(default=1, ge=1),
    db: Session = Depends(get_db),  # noqa: B008
) -> HTMLResponse:
    chosen = filters(months, date_from, date_to, hide_maybe_listed, news)
    # No filters in the address: the last ones chosen, not «all the time».
    if not any(name in request.query_params for name in FILTER_NAMES):
        chosen = remembered(request.cookies.get(FILTERS_COOKIE, "")) or chosen
    found, total, maybe, news_counts = _rows(db, chosen)
    on_page = _details(db, found[(page - 1) * PAGE_SIZE : page * PAGE_SIZE])
    rows = "".join(
        _html_row(position, row)
        for position, row in enumerate(on_page, start=(page - 1) * PAGE_SIZE + 1)
    )
    keep = chosen.query()
    # A period button clears the dates: the months are the choice then.
    periods = " ".join(
        f'<button type="submit" name="months" value="{key}" '
        f'class="chip{" active" if key == chosen.months and not chosen.custom else ""}" '
        "onclick=\"this.form.date_from.value='';this.form.date_to.value=''\">"
        f"{label}</button>"
        for key, label in PERIODS.items()
    )
    pages = (total + PAGE_SIZE - 1) // PAGE_SIZE
    pages_html = pager("/ui/political", keep, page, pages)
    news_options = "".join(
        f'<option value="{key}"{" selected" if key == chosen.news else ""}>{escape(label)}'
        f"{f' ({news_counts.get(key, 0)})' if key != 'all' else ''}</option>"
        for key, label in NEWS_FILTERS.items()
        if key in ("all", NEW_CASE, SENTENCE, ONGOING) or news_counts.get(key)
    )
    dates = (
        _date_field("date_from", "с", chosen.date_from)
        + _date_field("date_to", "по", chosen.date_to)
        + '<button type="submit">Показать</button>'
    )
    # The chosen months ride along hidden; a period button, sent later, wins over them.
    body = f"""<form method="get" action="/ui/political" class="toolbar" id="political-filters">
  <input type="hidden" name="months" value="{chosen.months}">
  <span class="chips">{periods}</span>
  <span class="chips">{dates}</span>
  <div class="filter-row">
    <!-- Unticked, only the hidden «false» is sent; ticked, the box's «true» comes last. -->
    <input type="hidden" name="hide_maybe_listed" value="false">
    <label class="check" title="Совпали имя и фамилия с перечнем, отчества нет с одной из сторон">
      <input id="box-hide-maybe" type="checkbox" name="hide_maybe_listed" value="true"{
        " checked" if chosen.hide_maybe_listed else ""
    } onchange="this.form.submit()"> Скрыть возможных в перечне ({maybe})</label>
    <!-- The form's own values, not the page's: dates picked but not shown yet count. -->
    <label class="field">Свежая новость <select name="news" onchange="this.form.submit()">{
        news_options
    }</select></label>
    <button type="submit" class="secondary" formaction="/ui/political/export.xlsx">Скачать Excel</button>
  </div>
</form>
{funnel_line(funnel(db))}
<p class="muted">Найдено: {total}. Фигуранты уголовных дел, дело которых — политическое
преследование. Перечень Росфинмониторинга подтверждает личность: дата рождения и место из
него — рядом с именем; «возможно в перечне» — совпали только имя и фамилия, может быть тёзка.
Жирная статья — из списка политических; «Последняя новость» показывает, свежий ли случай.</p>
<table><thead><tr><th>№</th><th>Фамилия Имя</th><th>Свежая новость</th><th>Регион</th><th>Перечень РФМ</th><th>Статьи УК</th>
<th>Почему политическое</th><th>Мемориал</th><th>Первая новость</th><th>Последняя новость</th>
<th>Публикации</th></tr></thead><tbody>{rows}</tbody></table>
{pages_html}
<script>
// Dates picked but not shown yet are kept too: a reload shows them.
document.getElementById("political-filters").addEventListener("change", (event) => {{
  if (!("date" in event.target.dataset)) return;
  const form = new URLSearchParams(new FormData(event.target.form));
  document.cookie = "{FILTERS_COOKIE}=" + encodeURIComponent(form.toString()) +
    "; path=/ui/political; max-age=31536000; samesite=lax";
}});
</script>"""
    response = _page(
        "Результат",
        body,
        active="political",
        instruction=(
            "Люди, против которых заведены политические уголовные дела; перечень "
            "Росфинмониторинга подтверждает их личность."
        ),
        next_action=(
            "Выберите период, чтобы увидеть свежие случаи; скачайте Excel. Результат обновляют "
            "шаги 4–6 в «Управлении»."
        ),
        db=db,
    )
    response.set_cookie(
        FILTERS_COOKIE,
        urlencode(chosen.query()),
        max_age=365 * 24 * 3600,
        path="/ui/political",
        samesite="lax",
    )
    return response


def political_xlsx(rows: list[ListRow]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Результат"
    sheet.append(
        [
            "№",
            "Фамилия Имя",
            "Регион",
            "Статьи УК",
            "Почему политическое",
            "Мемориал",
            "Первая новость",
            "Последняя новость",
            "Перечень РФМ",
            "Публикации",
            "Свежая новость",
        ]
    )
    for position, row in enumerate(rows, start=1):
        link = next(
            (url for _, url, _ in row.links if url.startswith(("http://", "https://"))), None
        )
        values: list[Any] = [
            position,
            display_name(row.entity.name),
            _regions_text(row.entity) or None,
            _article_text(row.articles) or None,
            _basis(row),
            row.memorial,
            row.first_published.replace(tzinfo=None) if row.first_published else None,
            row.entity.last_published_at.replace(tzinfo=None)
            if row.entity.last_published_at
            else None,
            rf_text(row) or None,
            "\n".join(url for _, url, _ in row.links) or None,
            KIND_LABELS.get(row.news_kind, row.news_kind) if row.news_kind else None,
        ]
        sheet.append(values)
        line = position + 1
        # Names, reasons and URLs come from scraped sources: never let «=» become a formula.
        for column in (2, 3, 4, 5, 6, 10):
            if sheet.cell(row=line, column=column).value is not None:
                sheet.cell(row=line, column=column).data_type = "s"
        for column in (7, 8):
            sheet.cell(row=line, column=column).number_format = "DD.MM.YYYY"
        if link is not None:
            sheet.cell(row=line, column=10).hyperlink = link
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_name(
    chosen: Filters, found: list[tuple[EntityGroupRecord, EntityGroupPoliticsRecord]], today: date
) -> str:
    """`result_<from>_<to>.xlsx`: the dates chosen, else the period's start, else the
    earliest latest news of the rows; up to the date chosen, else today."""
    start = chosen.date_from
    if start is None and chosen.months:
        start = today - timedelta(days=30 * chosen.months)
    if start is None:
        dates = [row.last_published_at for row, _ in found if row.last_published_at is not None]
        start = min(dates).date() if dates else today
    end = chosen.date_to or today
    return f"result_{start.isoformat()}_{end.isoformat()}.xlsx"


@router.get("/ui/political/export.xlsx")
def ui_political_export(
    months: int = Query(default=0),
    date_from: str = Query(default="", max_length=10),
    date_to: str = Query(default="", max_length=10),
    hide_maybe_listed: bool = Query(default=False),
    news: str = Query(default="all", max_length=16),
    db: Session = Depends(get_db),  # noqa: B008
) -> Response:
    """Every row of the page's filters, not only one page."""
    chosen = filters(months, date_from, date_to, hide_maybe_listed, news)
    found, _, _, _ = _rows(db, chosen)
    name = export_name(chosen, found, datetime.now(UTC).date())
    return Response(
        political_xlsx(_details(db, found)),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )
