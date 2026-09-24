"""The candidates export as a file: Excel, built from rows, not requests.

The route in `web.ui.candidates` chooses the rows and wraps the bytes in a response.
"""

from collections.abc import Sequence
from io import BytesIO

from openpyxl import Workbook

from web.candidate_rows import _CANDIDATE_CATEGORIES, _CandidateRow, _news_day, _surname_first


def candidates_xlsx(rows: Sequence[_CandidateRow]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Кандидаты"
    sheet.append(["№", "Фамилия Имя", "Дата новости", "Категория", "Причины", "Ссылка"])
    for position, (candidate, news) in enumerate(rows, start=1):
        link = news.url if news is not None else None
        published = _news_day(news.published_at) if news is not None and news.published_at else None
        category = (
            _CANDIDATE_CATEGORIES.get(news.event_type, news.event_type)
            if news is not None and news.event_type
            else None
        )
        reasons = "; ".join(candidate.persecution_reasons) or None
        sheet.append(
            [position, _surname_first(candidate.canonical_name), published, category, reasons, link]
        )
        row = position + 1
        # Names and URLs come from scraped sources: never let a leading "=" become a formula.
        for column, value in ((2, True), (5, reasons), (6, link)):
            if value is not None:
                sheet.cell(row=row, column=column).data_type = "s"
        if published is not None:
            sheet.cell(row=row, column=3).number_format = "DD.MM.YYYY"
        # Only web links are clickable: a scraped «javascript:» URL stays plain text.
        if link is not None and link.startswith(("http://", "https://")):
            sheet.cell(row=row, column=6).hyperlink = link
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
