"""Excel representation of the same source reports displayed by /ui/cases."""

import re
from collections.abc import Sequence
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from channel_feed.case_age import AGE_LABELS
from channel_feed.case_reports import CaseReport
from web.candidate_rows import _CANDIDATE_CATEGORIES

_XML_CONTROLS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def case_reports_xlsx(reports: Sequence[CaseReport]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Дела и события"
    sheet.append(
        [
            "Давность дела",
            "Фигурант",
            "Дата новости",
            "Событие",
            "Время в цитате",
            "Почему такой статус",
            "Основание отбора",
            "Цитата",
            "Источник",
            "Ссылка",
        ]
    )
    for report in reports:
        sheet.append(
            [
                AGE_LABELS[report.timing.age],
                "; ".join(report.names) or "Имя не определено",
                report.published,
                _CANDIDATE_CATEGORIES.get(report.event_type, report.event_type),
                report.timing.evidence,
                report.timing.reason,
                report.basis,
                _XML_CONTROLS.sub("", report.quotation),
                _XML_CONTROLS.sub("", report.source),
                _XML_CONTROLS.sub("", report.url),
            ]
        )
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.cell(sheet.max_row, 3).number_format = "DD.MM.YYYY"
        if report.url.startswith(("https://", "http://")):
            sheet.cell(sheet.max_row, 10).hyperlink = report.url
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="254B60")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column, width in enumerate((25, 30, 15, 20, 22, 44, 44, 90, 25, 45), 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
