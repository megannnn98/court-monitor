"""The bot's `/people` and `/export` are the candidates of `/ui/candidates`, on PostgreSQL."""

from __future__ import annotations

import asyncio
import re
from collections.abc import Callable
from datetime import UTC, date, datetime
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy.orm import Session, sessionmaker
from support.research_db_fixtures import ResearchSeeder

from db.orm_models import RosfinmonitoringSnapshotRecord
from telegram_bot.authorization import Authorization
from telegram_bot.candidates import CandidatesRepository
from telegram_bot.config import TelegramBotSettings
from telegram_bot.handlers import CommandHandlers
from telegram_bot.people_service import PeopleService
from web.candidate_rows import select_candidate_rows

ALLOWED = 42
SETTINGS = TelegramBotSettings.from_env(
    {"TELEGRAM_BOT_TOKEN": "123:secret", "TELEGRAM_ALLOWED_USER_IDS": str(ALLOWED)}
)
PERIOD = ("2026-09-01", "2026-09-19")
IN_PERIOD = datetime(2026, 9, 10, 9, tzinfo=UTC)
AFTER_PERIOD = datetime(2026, 9, 21, 9, tzinfo=UTC)


class _NoUpdates:
    def start_update(self) -> None:
        raise AssertionError("not used")

    def last_update(self) -> None:
        raise AssertionError("not used")


def _handlers(session_factory: sessionmaker[Session]) -> CommandHandlers:
    async def run_here[T](work: Callable[[], T]) -> T:
        return work()

    return CommandHandlers(
        settings=SETTINGS,
        authorization=Authorization(SETTINGS),
        people=PeopleService(CandidatesRepository(session_factory), SETTINGS),
        updates=_NoUpdates(),  # type: ignore[arg-type]
        run_blocking=run_here,
    )


def _person_in_the_news(
    seed: ResearchSeeder,
    name: str,
    *,
    published_at: datetime = IN_PERIOD,
    classification: str | None = "political",
    rf_status: str | None = "not_matched",
    snapshot_id: int,
) -> int:
    person_id = seed.person(name)
    if classification is not None:
        seed.classification(
            person_id, classification, 0.9, reasons=["Политическая статья: УК РФ ст. 207.3"]
        )
    if rf_status is not None:
        seed.match(person_id, snapshot_id, rf_status, 0.8)
    source_id = seed.source(f"source-{person_id}", f"https://news-{person_id}.example.test")
    _, run_id = seed.article(
        source_id,
        external_id=f"news-{person_id}",
        title="Новость",
        text=f"Суд арестовал {name}.",
        published_at=published_at,
    )
    seed.event(
        run_id,
        f"Суд арестовал {name}",
        event_type="arrest",
        event_date=published_at,
        links=[(person_id, "subject")],
    )
    return person_id


def _seed(session_factory: sessionmaker[Session]) -> tuple[int, dict[str, int]]:
    with session_factory() as session:
        seed = ResearchSeeder(session)
        older = seed.snapshot("older-snapshot")
        snapshot_id = seed.snapshot("newest-snapshot")
        session.get_one(RosfinmonitoringSnapshotRecord, older).snapshot_date = datetime(
            2026, 8, 1, tzinfo=UTC
        )
        session.get_one(RosfinmonitoringSnapshotRecord, snapshot_id).snapshot_date = datetime(
            2026, 9, 1, tzinfo=UTC
        )
        people = {
            "candidate": _person_in_the_news(seed, "Иван Кандидатов", snapshot_id=snapshot_id),
            "second": _person_in_the_news(seed, "Петр Вторых", snapshot_id=snapshot_id),
            # In the news of the period, but not `political + not_matched`:
            "on_the_list": _person_in_the_news(
                seed, "Анна Вперечне", rf_status="matched", snapshot_id=snapshot_id
            ),
            "not_political": _person_in_the_news(
                seed, "Олег Неполитический", classification="non_political", snapshot_id=snapshot_id
            ),
            "never_checked": _person_in_the_news(
                seed, "Яков Непроверенный", rf_status=None, snapshot_id=snapshot_id
            ),
            "only_in_the_news": _person_in_the_news(
                seed,
                "Мария Просто",
                classification=None,
                rf_status=None,
                snapshot_id=snapshot_id,
            ),
            # A candidate, but the news is after the period's end.
            "later": _person_in_the_news(
                seed, "Семен Позже", published_at=AFTER_PERIOD, snapshot_id=snapshot_id
            ),
            # `not_matched` only in the older snapshot: not a candidate of the newest one.
            "older_snapshot_only": _person_in_the_news(
                seed, "Глеб Старый", rf_status=None, snapshot_id=snapshot_id
            ),
        }
        seed.match(people["older_snapshot_only"], older, "not_matched", 0.8)
        session.commit()
    return snapshot_id, people


def test_people_and_export_are_one_cohort_of_page_candidates(
    session_factory: sessionmaker[Session],
) -> None:
    snapshot_id, people = _seed(session_factory)
    bot = _handlers(session_factory)

    listed = asyncio.run(bot.handle("people", list(PERIOD), ALLOWED))
    exported = asyncio.run(bot.handle("export", list(PERIOD), ALLOWED))

    message = "\n".join(listed.messages)
    assert exported.document is not None
    sheet = load_workbook(BytesIO(exported.document.content)).active
    assert sheet is not None
    exported_names = [row[1] for row in list(sheet.iter_rows(values_only=True))[1:]]
    # The same people, in the same order, in the chat and in the file.
    listed_names = re.findall(r"\d+\. <b>(.*?)</b>", message)
    assert listed_names == exported_names == ["Кандидатов Иван", "Вторых Петр"]
    for outsider in (
        "Вперечне",
        "Неполитический",
        "Непроверенный",
        "Просто",
        "Позже",
        "Старый",
    ):
        assert outsider not in message

    # And it is the page's selection for that period and the newest snapshot.
    with session_factory() as session:
        page_rows = select_candidate_rows(
            session,
            snapshot_id=snapshot_id,
            period_start=date(2026, 9, 1),
            period_end=date(2026, 9, 19),
        )
    assert [row.candidate.person_id for row in page_rows] == [
        people["candidate"],
        people["second"],
    ]


def test_the_period_end_is_inclusive_and_excludes_later_news(
    session_factory: sessionmaker[Session],
) -> None:
    _seed(session_factory)
    bot = _handlers(session_factory)

    through_the_later_news = asyncio.run(
        bot.handle("people", ["2026-09-01", "2026-09-21"], ALLOWED)
    )

    assert "Позже Семен" in "\n".join(through_the_later_news.messages)
