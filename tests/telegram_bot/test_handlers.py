"""Handlers with fake services: answers, authorization, errors — no Telegram, no database."""

from __future__ import annotations

import asyncio
import logging
from collections.abc import Callable, Sequence
from datetime import UTC, date, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from candidates.models import PoliticalPersecutionCandidate, RosfinmonitoringStatus
from monitoring.models import FailureKind
from operator_console import (
    OPERATION_DEFINITIONS,
    OperationParameters,
    OperationRun,
    OperationRunStatus,
)
from telegram_bot import formatting
from telegram_bot.app import _parse
from telegram_bot.authorization import Authorization
from telegram_bot.candidates import CandidatesResult
from telegram_bot.config import TelegramBotSettings
from telegram_bot.handlers import Answer, CommandHandlers
from telegram_bot.people_service import PeopleQuery, PeopleQueryError, parse_people_query
from telegram_bot.period_keyboard import Button
from telegram_bot.update_service import (
    RunFailures,
    UpdateAlreadyRunning,
    UpdateStarted,
    UpdateStatus,
)
from web.candidate_rows import CandidateRow, _CandidateNews

ALLOWED = 42
DENIED = 4242
SETTINGS = TelegramBotSettings.from_env(
    {"TELEGRAM_BOT_TOKEN": "123:secret", "TELEGRAM_ALLOWED_USER_IDS": str(ALLOWED)}
)
RUN = OperationRun(
    id=42,
    operation=OPERATION_DEFINITIONS["monitor"],
    parameters=OperationParameters(limit=50),
    status=OperationRunStatus.RUNNING,
    created_at=datetime(2026, 9, 19, 14, 32, tzinfo=UTC),
    started_at=datetime(2026, 9, 19, 14, 32, tzinfo=UTC),
)
PERIOD = (date(2026, 9, 1), date(2026, 9, 19))
EMPTY_RESULT = CandidatesResult(date_from=PERIOD[0], date_to=PERIOD[1], snapshot_id=1, rows=[])


def candidate_row(
    person_id: int,
    name: str,
    *,
    url: str = "https://news.example/1",
    event_type: str | None = "arrest",
) -> CandidateRow:
    return CandidateRow(
        PoliticalPersecutionCandidate(
            person_id=person_id,
            canonical_name=name,
            normalized_name=name.lower(),
            persecution_status="political",
            persecution_confidence=0.9,
            persecution_reasons=["Политическая статья: УК РФ ст. 207.3"],
            rosfinmonitoring_status=RosfinmonitoringStatus.NOT_MATCHED,
            rosfinmonitoring_match_confidence=0.8,
            event_count=1,
            alias_count=0,
            last_event_date=None,
        ),
        _CandidateNews(url, datetime(2026, 9, 18, 10, tzinfo=UTC), event_type),
    )


def candidates(count: int) -> CandidatesResult:
    return CandidatesResult(
        date_from=PERIOD[0],
        date_to=PERIOD[1],
        snapshot_id=1,
        rows=[candidate_row(index, f"Иван Человек{index}") for index in range(count)],
    )


class FakeUpdates:
    def __init__(
        self,
        outcome: UpdateStarted | UpdateAlreadyRunning | None = None,
        status: UpdateStatus | None = None,
    ) -> None:
        self.outcome = outcome or UpdateStarted(run=RUN, sources=("ovd-info", "sota-vision"))
        self.status = status
        self.started = 0

    def start_update(self) -> UpdateStarted | UpdateAlreadyRunning:
        self.started += 1
        return self.outcome

    def last_update(self) -> UpdateStatus | None:
        return self.status


class FakePeople:
    def __init__(self, result: CandidatesResult = EMPTY_RESULT) -> None:
        self.result = result
        self.queries: list[PeopleQuery] = []

    def parse(self, arguments: Sequence[str]) -> PeopleQuery:
        return parse_people_query(
            arguments,
            timezone=SETTINGS.timezone,
            default_limit=SETTINGS.people_default_limit,
            max_limit=SETTINGS.people_max_limit,
        )

    def parse_export(self, arguments: Sequence[str]) -> PeopleQuery:
        if len(arguments) != 2:
            raise PeopleQueryError("нужны две даты")
        return self.parse(arguments)

    def period_query(self, date_from: date, date_to: date) -> PeopleQuery:
        return parse_people_query(
            [date_from.isoformat(), date_to.isoformat()],
            timezone=SETTINGS.timezone,
            default_limit=SETTINGS.people_default_limit,
            max_limit=SETTINGS.people_max_limit,
        )

    def people(self, query: PeopleQuery) -> CandidatesResult:
        self.queries.append(query)
        return self.result

    def export(self, query: PeopleQuery) -> CandidatesResult:
        self.queries.append(query)
        return self.result


def handlers(
    *, updates: FakeUpdates | None = None, people: FakePeople | None = None
) -> CommandHandlers:
    async def run_here[T](work: Callable[[], T]) -> T:
        # The "blocking" call runs inline: no test depends on thread timing.
        return work()

    return CommandHandlers(
        settings=SETTINGS,
        authorization=Authorization(SETTINGS),
        people=people or FakePeople(),  # type: ignore[arg-type]
        updates=updates or FakeUpdates(),  # type: ignore[arg-type]
        run_blocking=run_here,
    )


def answer(
    command: str,
    *arguments: str,
    user_id: int | None = ALLOWED,
    updates: FakeUpdates | None = None,
    people: FakePeople | None = None,
) -> Answer:
    bot = handlers(updates=updates, people=people)
    return asyncio.run(bot.handle(command, list(arguments), user_id))


def reply(
    command: str,
    *arguments: str,
    user_id: int | None = ALLOWED,
    updates: FakeUpdates | None = None,
    people: FakePeople | None = None,
) -> list[str]:
    return answer(command, *arguments, user_id=user_id, updates=updates, people=people).messages


def test_start_lists_only_the_candidate_commands() -> None:
    messages = reply("start")

    assert messages == [formatting.START]
    commands = {word for word in messages[0].split() if word.startswith("/")}
    assert commands == {"/people", "/export", "/update", "/status", "/help"}
    assert "кандидат" in messages[0].lower()
    assert "новост" not in messages[0].lower()


def test_help_describes_the_candidate_cohort_and_the_update_commands() -> None:
    message = reply("help")[0]

    assert "Росфинмониторинга" in message and "«Кандидаты»" in message
    assert "по московскому времени" in message
    assert "/update — обновить данные кандидатов" in message
    assert "/status" in message and "YYYY-MM-DD" in message


def test_update_answers_with_the_run_id() -> None:
    updates = FakeUpdates()

    message = reply("update", updates=updates)[0]

    assert "Обновление запущено" in message
    assert "Run: #42" in message
    assert updates.started == 1


def test_update_reports_a_run_already_going_on() -> None:
    updates = FakeUpdates(outcome=UpdateAlreadyRunning(run=RUN))

    message = reply("update", updates=updates)[0]

    assert "Обновление уже выполняется" in message
    assert "Run: #42" in message


def test_status_without_any_run() -> None:
    assert "ещё ни разу не запускалось" in reply("status")[0]


def test_status_shows_counters_and_failures() -> None:
    status = UpdateStatus(
        run=RUN,
        sources_done=2,
        sources_total=2,
        current_source=None,
        derived_done=True,
        documents_discovered=34,
        documents_ingested=12,
        articles_extracted=12,
        events_created=27,
        persons_created=8,
        persons_linked=5,
        reviews_created=3,
        error_count=1,
        failures=(
            RunFailures(
                stage="ingestion", count=1, kind=FailureKind.RETRYABLE, message="ReadTimeout"
            ),
        ),
    )

    message = reply("status", updates=FakeUpdates(status=status))[0]

    assert "Run #42" in message
    assert "Обнаружено публикаций: 34" in message
    assert "Ошибок: 1" in message
    assert "ingestion: 1 × ReadTimeout (повторяемая)" in message


def test_people_reports_an_empty_period() -> None:
    message = reply("people", "2026-09-01", "2026-09-19")[0]

    assert "Найдено всего: 0" in message
    assert "не найдено" in message


def test_people_lists_candidates_in_the_page_order() -> None:
    result = CandidatesResult(
        date_from=PERIOD[0],
        date_to=PERIOD[1],
        snapshot_id=1,
        rows=[
            candidate_row(1, "Иван <Иванов>", url="https://news.example/7?a=1&b=2"),
            candidate_row(2, "Петр Петров", event_type="sentence"),
        ],
    )

    message = reply("people", "2026-09-01", "2026-09-19", people=FakePeople(result))[0]

    assert "Кандидаты 2026-09-01 — 2026-09-19" in message
    assert "Найдено всего: 2" in message and "Показано: 2" in message
    # Names are written surname first, as on the page; text is escaped, links stay links.
    assert "1. <b>&lt;Иванов&gt; Иван</b>" in message
    assert "18.09.2026 · Арест" in message
    assert '<a href="https://news.example/7?a=1&amp;b=2">Новость</a>' in message
    assert message.index("Иван") < message.index("Петров Петр")


def test_people_shows_at_most_the_limit_but_counts_all() -> None:
    message = reply("people", "2026-09-01", "2026-09-19", "2", people=FakePeople(candidates(5)))[0]

    assert "Найдено всего: 5" in message and "Показано: 2" in message


def test_without_a_snapshot_there_are_no_candidates() -> None:
    result = CandidatesResult(date_from=PERIOD[0], date_to=PERIOD[1], snapshot_id=None, rows=[])

    people = reply("people", "2026-09-01", "2026-09-19", people=FakePeople(result))
    exported = answer("export", "2026-09-01", "2026-09-19", people=FakePeople(result))

    assert people == [formatting.NO_SNAPSHOT]
    assert exported.document is None and exported.messages == [formatting.NO_SNAPSHOT]


def test_people_passes_the_parsed_period_to_the_service() -> None:
    people = FakePeople()

    reply("people", "2026-09-01", "2026-09-19", "100", people=people)

    assert people.queries[0].limit == 100
    assert people.queries[0].date_from.isoformat() == "2026-09-01"


@pytest.mark.parametrize(
    "arguments",
    [
        ("2026-09-01",),
        ("01.09.2026", "19.09.2026"),
        ("2026-99-99", "2026-09-19"),
        ("2026-09-20", "2026-09-01"),
        ("2026-09-01", "2026-09-19", "abc"),
    ],
)
def test_people_explains_the_format_on_a_bad_argument(arguments: tuple[str, ...]) -> None:
    message = reply("people", *arguments)[0]

    assert "/people YYYY-MM-DD YYYY-MM-DD [limit]" in message
    assert "Traceback" not in message


def test_an_unknown_command_is_answered_shortly() -> None:
    assert reply("dance") == [formatting.UNKNOWN_COMMAND]


def test_an_unauthorized_user_gets_nothing_done() -> None:
    updates = FakeUpdates()

    assert reply("update", user_id=DENIED, updates=updates) == [formatting.NOT_AUTHORIZED]
    assert updates.started == 0


def test_a_message_without_a_user_is_denied() -> None:
    assert reply("people", "2026-09-01", "2026-09-19", user_id=None) == [formatting.NOT_AUTHORIZED]


def test_an_unexpected_error_never_reaches_the_user(caplog: pytest.LogCaptureFixture) -> None:
    class Exploding(FakeUpdates):
        def start_update(self) -> UpdateStarted | UpdateAlreadyRunning:
            raise RuntimeError("postgresql://user:password@host/db is unreachable")

    with caplog.at_level(logging.ERROR, logger="telegram_bot.handlers"):
        messages = reply("update", updates=Exploding())

    assert messages == [formatting.UNEXPECTED_ERROR]
    assert "password" not in messages[0]
    # The server log keeps the traceback.
    assert caplog.records and caplog.records[-1].exc_info is not None


def test_a_command_addressed_to_the_bot_is_understood() -> None:
    assert _parse("/people@court_monitor_bot 2026-09-01 2026-09-19") == (
        "people",
        ["2026-09-01", "2026-09-19"],
    )


def test_plain_text_is_not_a_command() -> None:
    assert _parse("привет") == ("", [])


def test_a_very_long_name_stays_inside_one_message_with_its_link() -> None:
    result = CandidatesResult(
        date_from=PERIOD[0],
        date_to=PERIOD[1],
        snapshot_id=1,
        rows=[
            candidate_row(
                index,
                f"Человек{index} " + "Длинноимённый" * 30,
                url=f"https://news.example/{index}",
            )
            for index in range(50)
        ],
    )

    messages = reply("people", "2026-09-01", "2026-09-19", people=FakePeople(result))

    for message in messages:
        assert len(message) <= 4096
        # Every link and every bold name opened in a part is closed in the same part.
        assert message.count("<a href=") == message.count("</a>")
        assert message.count("<b>") == message.count("</b>")


def test_export_answers_with_the_page_spreadsheet() -> None:
    people = FakePeople(candidates(3))

    result = answer("export", "2026-09-01", "2026-09-19", people=people)

    assert result.document is not None
    assert result.document.filename == "candidates-2026-09-01-2026-09-19.xlsx"
    # The file of the page's «Export to Excel»: its header, one row per candidate.
    sheet = load_workbook(BytesIO(result.document.content)).active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("№", "Фамилия Имя", "Дата новости", "Категория", "Причины", "Ссылка")
    assert [row[1] for row in rows[1:]] == ["Человек0 Иван", "Человек1 Иван", "Человек2 Иван"]
    assert "Кандидатов в файле: 3" in result.messages[0]


def test_export_of_an_empty_period_sends_no_file() -> None:
    result = answer("export", "2026-09-01", "2026-09-19")

    assert result.document is None
    assert "выгружать нечего" in result.messages[0]


@pytest.mark.parametrize(
    "arguments",
    [("2026-09-01",), ("01.09.2026", "19.09.2026"), ("2026-09-01", "2026-09-19", "100")],
)
def test_export_explains_its_format(arguments: tuple[str, ...]) -> None:
    result = answer("export", *arguments)

    assert result.document is None
    assert "/export YYYY-MM-DD YYYY-MM-DD" in result.messages[0]


def test_export_is_closed_by_the_allowlist() -> None:
    people = FakePeople(candidates(3))

    result = answer("export", "2026-09-01", "2026-09-19", user_id=DENIED, people=people)

    assert result.document is None
    assert result.messages == [formatting.NOT_AUTHORIZED]
    assert people.queries == []


def press(data: str, *, user_id: int | None = ALLOWED, people: FakePeople | None = None) -> Answer:
    bot = handlers(people=people)
    return asyncio.run(bot.handle_callback(data, user_id))


def pressable(keyboard: list[list[Button]]) -> list[Button]:
    return [button for row in keyboard for button in row if button.data]


def test_people_without_arguments_offers_the_period() -> None:
    result = answer("people")

    assert result.keyboard is not None
    assert "Период кандидатов" in result.messages[0]
    assert any(button.text == "7 дней" for button in pressable(result.keyboard))


def test_export_without_arguments_offers_the_period() -> None:
    result = answer("export")

    assert result.keyboard is not None
    assert "Период выгрузки кандидатов" in result.messages[0]


def test_a_preset_button_runs_the_search() -> None:
    people = FakePeople()
    keyboard = answer("people").keyboard
    assert keyboard is not None
    seven_days = next(button for button in pressable(keyboard) if button.text == "7 дней")

    result = press(seven_days.data, people=people)

    assert "Найдено всего" in result.messages[0]
    assert people.queries and (people.queries[0].date_to - people.queries[0].date_from).days == 6


def test_a_preset_button_of_export_sends_the_file() -> None:
    people = FakePeople(candidates(2))
    keyboard = answer("export").keyboard
    assert keyboard is not None
    month = next(button for button in pressable(keyboard) if button.text.startswith("Этот месяц"))

    result = press(month.data, people=people)

    assert result.document is not None
    assert result.document.filename.endswith(".xlsx")


def test_the_calendar_button_opens_a_month_grid() -> None:
    keyboard = answer("people").keyboard
    assert keyboard is not None
    calendar_button = next(
        button for button in pressable(keyboard) if button.text == "📅 Другой период"
    )

    result = press(calendar_button.data)

    assert result.edit
    assert result.keyboard is not None
    assert (
        "выберите начало" in result.messages[0] or "выберите начало" in result.keyboard[0][0].text
    )


def test_two_day_presses_run_the_search() -> None:
    people = FakePeople()
    grid = press(
        next(
            button
            for button in pressable(answer("people").keyboard or [])
            if button.text == "📅 Другой период"
        ).data
    ).keyboard
    assert grid is not None
    first = next(button for button in pressable(grid) if ":day:" in button.data)

    after_first = press(first.data, people=people)
    assert after_first.keyboard is not None
    assert people.queries == []

    second = [button for button in pressable(after_first.keyboard) if ":day:" in button.data][-1]
    result = press(second.data, people=people)

    assert "Найдено всего" in result.messages[0]
    assert people.queries


def test_a_label_press_changes_nothing() -> None:
    result = press("p:noop")

    assert result.messages == []
    assert result.keyboard is None


def test_cancel_closes_the_chooser() -> None:
    result = press("p:cancel:people:")

    assert result.edit
    assert "отменён" in result.messages[0]


def test_an_unreadable_button_asks_to_repeat_the_command() -> None:
    result = press("garbage")

    assert result.edit
    assert "устарела" in result.messages[0]


def test_a_button_press_is_closed_by_the_allowlist() -> None:
    people = FakePeople()

    result = press("p:preset:people:7d", user_id=DENIED, people=people)

    assert result.messages == [formatting.NOT_AUTHORIZED]
    assert people.queries == []
