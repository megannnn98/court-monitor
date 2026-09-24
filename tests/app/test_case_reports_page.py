"""Page/export agreement without a Person, RF snapshot or guessed identity."""

from datetime import date
from io import BytesIO
from unittest.mock import Mock

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api import app, get_db
from channel_feed.case_age import CaseAge, CaseTiming
from channel_feed.case_reports import CaseReport
from web.case_export import case_reports_xlsx
from web.ui import cases, layout


def _report() -> CaseReport:
    return CaseReport(
        event_id=1,
        article_id=1,
        published=date(2026, 9, 23),
        event_type="case_opened",
        names=(),
        timing=CaseTiming(CaseAge.NEW, "Есть дата возбуждения", "сегодня"),
        basis="УК: 275",
        quotation="Сегодня возбудили дело по ст. 275 УК РФ.",
        title="Имя не сообщается",
        source="Источник",
        url="https://example.test/1",
    )


def test_page_and_excel_use_identical_filters_and_keep_unnamed(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    loader = Mock(return_value=[_report()])
    monkeypatch.setattr(cases, "load_case_reports", loader)
    monkeypatch.setattr(
        layout,
        "_status_counts",
        lambda db: {
            "articles": 1,
            "persons": 0,
            "pending_reviews": 0,
            "latest_run": "нет",
        },
    )
    db = Mock()
    app.dependency_overrides[get_db] = lambda: db
    filters = "date_from=2026-08-01&date_to=2026-09-23&age=new&unnamed_only=1&limit=10"
    try:
        with TestClient(app) as client:
            page = client.get(f"/ui/cases?{filters}")
            exported = client.get(f"/ui/cases/export.xlsx?{filters}")
    finally:
        app.dependency_overrides.pop(get_db, None)
    assert page.status_code == exported.status_code == 200
    assert "Имя не определено" in page.text
    assert "Сегодня возбудили дело" in page.text
    assert "Найдено сообщений: 1" in page.text
    assert loader.call_args_list[0] == loader.call_args_list[1]
    assert loader.call_args.kwargs == {
        "period_start": date(2026, 8, 1),
        "period_end": date(2026, 9, 23),
        "age": CaseAge.NEW,
        "unnamed_only": True,
    }
    sheet = load_workbook(BytesIO(exported.content)).active
    assert sheet is not None
    assert sheet["A2"].value == "Новое дело"
    assert sheet["B2"].value == "Имя не определено"
    assert sheet["H2"].value == _report().quotation


@pytest.mark.parametrize("path", ["/ui/cases", "/ui/cases/export.xlsx"])
@pytest.mark.parametrize(
    "filters",
    [
        "date_from=2026-09-23&date_to=2026-08-01",
        "date_from=2020-01-01&date_to=2026-09-23",
        "age=invalid",
        "limit=0",
    ],
)
def test_invalid_filters_are_rejected(path: str, filters: str) -> None:
    app.dependency_overrides[get_db] = lambda: Mock()
    try:
        with TestClient(app) as client:
            assert client.get(f"{path}?{filters}").status_code == 422
    finally:
        app.dependency_overrides.pop(get_db, None)


def test_export_preserves_dates_and_treats_source_text_as_text() -> None:
    from dataclasses import replace

    report = replace(_report(), names=("=1+1",), quotation="=1+1", url="javascript:alert(1)")
    sheet = load_workbook(BytesIO(case_reports_xlsx([report]))).active
    assert sheet is not None
    assert sheet["B2"].data_type == sheet["H2"].data_type == "s"
    assert sheet["C2"].is_date
    assert sheet["J2"].hyperlink is None
    assert sheet.freeze_panes == "C2"
    assert sheet.auto_filter.ref == "A1:J2"
