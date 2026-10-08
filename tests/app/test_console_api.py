"""The read routes behind the React console's «Все люди», «Публикации» and «Результат».

Each is read through the same functions as its legacy page, so these tests check the
answer against the legacy page's own rows and counts."""

from __future__ import annotations

from collections.abc import Iterator
from contextlib import contextmanager
from datetime import UTC, datetime, timedelta
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session, sessionmaker
from support.db_fixtures import DatabaseSeeder

from db.orm_models import (
    EntityGroupPoliticsRecord,
    EntityGroupRecord,
    EntityGroupRfMatchRecord,
    EntityMentionRecord,
    RosfinmonitoringEntryRecord,
    RosfinmonitoringSnapshotRecord,
)
from entities.collector import EntityCollector
from web.app import app
from web.dependencies import get_db


@contextmanager
def _client(session_factory: sessionmaker[Session]) -> Iterator[TestClient]:
    def override() -> Iterator[Session]:
        with session_factory() as session:
            yield session

    app.dependency_overrides[get_db] = override
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.pop(get_db, None)


def _person(
    session: Session, seed: DatabaseSeeder, run: int, surface: str, first: str, last: str
) -> int:
    mention_id = seed.mention(run, surface, person_id=None)
    session.get_one(EntityMentionRecord, mention_id).normalized_data = {
        "first_name": first,
        "last_name": last,
        "patronymic": None,
    }
    return mention_id


def _collected(session_factory: sessionmaker[Session]) -> tuple[int, int]:
    """Two publications naming Моор, one naming Иванов too; the people collected."""
    with session_factory() as session:
        seed = DatabaseSeeder(session)
        source = seed.source("news", "https://news.example.test")
        arrest, run = seed.article(
            source,
            external_id="arrest",
            title="Арест Моора",
            text="Суд арестовал Александра Моора и Ивана Иванова.",
            published_at=datetime(2026, 9, 1, tzinfo=UTC),
        )
        arrested = _person(session, seed, run, "Александра Моора", "Александр", "Моора")
        _person(session, seed, run, "Ивана Иванова", "Иван", "Иванов")
        seed.event(
            run,
            "Суд арестовал",
            event_type="arrest",
            event_date=None,
            links=[],
            entity_links=[(arrested, "subject")],
        )
        sentence, run = seed.article(
            source,
            external_id="sentence",
            title="Приговор Моору",
            text="Александру Моору вынесли приговор.",
            published_at=datetime(2026, 9, 10, tzinfo=UTC),
        )
        _person(session, seed, run, "Александру Моору", "Александр", "Моору")
        seed.event(run, "вынесли приговор", event_type="sentence", event_date=None, links=[])
        session.commit()
    EntityCollector(session_factory).run()
    return arrest, sentence


def test_the_people_are_the_legacy_list_s(session_factory: sessionmaker[Session]) -> None:
    _collected(session_factory)

    with _client(session_factory) as client:
        answer = client.get("/api/v1/entities", params={"role": "all"})
        found = client.get("/api/v1/entities", params={"role": "all", "q": "Моору"}).json()
        page = client.get("/ui/entities", params={"role": "all"}).text

    assert answer.status_code == 200
    body = answer.json()
    assert body["total"] == 2 and f"Найдено: {body['total']}" in page
    assert body["page"] == 1 and body["page_size"] == 100
    assert [row["name"] for row in body["items"]] == ["Моор Александр", "Иванов Иван"]
    for row in body["items"]:
        assert row["name"] in page
        assert row["dossier_url"].startswith("/ui/investigations/")
    moor = body["items"][0]
    assert (moor["mention_count"], moor["article_count"]) == (2, 2)
    assert {"kind": "arrest", "label": "Арест", "count": 1} in moor["events"]
    assert {option["value"] for option in body["roles"]} >= {"figurant", "all"}
    # Any form of the name finds the person, as on the legacy page.
    assert [row["name"] for row in found["items"]] == ["Моор Александр"]


def test_an_unknown_filter_is_refused(session_factory: sessionmaker[Session]) -> None:
    with _client(session_factory) as client:
        assert client.get("/api/v1/entities", params={"sort": "nonsense"}).status_code == 422


def test_the_publications_carry_their_people_and_events(
    session_factory: sessionmaker[Session],
) -> None:
    arrest, _ = _collected(session_factory)

    with _client(session_factory) as client:
        body = client.get("/api/v1/publications").json()
        found = client.get("/api/v1/publications", params={"q": "Приговор"}).json()
        page = client.get("/ui/publications").text

    assert body["total"] == 2 and f"Найдено: {body['total']}." in page
    assert [row["title"] for row in body["items"]] == ["Приговор Моору", "Арест Моора"]
    first = next(row for row in body["items"] if row["id"] == arrest)
    assert {person["name"] for person in first["people"]} == {"Моор Александр", "Иванов Иван"}
    assert first["events"] == [{"kind": "arrest", "label": "Арест", "count": 1}]
    assert body["sources"] == [{"id": body["sources"][0]["id"], "name": "news", "count": 2}]
    assert [row["title"] for row in found["items"]] == ["Приговор Моору"]


def test_an_article_s_people_and_events(session_factory: sessionmaker[Session]) -> None:
    arrest, _ = _collected(session_factory)

    with _client(session_factory) as client:
        mentions = client.get(f"/api/v1/articles/{arrest}/mentions").json()
        missing = client.get("/api/v1/articles/999999/mentions")
        page = client.get(f"/ui/articles/{arrest}").text

    for person in mentions["people"]:
        assert person["name"] in page
    assert {person["name"] for person in mentions["people"]} == {"Моор Александр", "Иванов Иван"}
    assert mentions["events"] == [{"kind": "arrest", "label": "Арест", "count": 1}]
    assert missing.status_code == 404 and missing.json() == {"detail": "Article not found"}


def _political(session_factory: sessionmaker[Session]) -> None:
    """Смирнова on the list with her patronymic, Иванов maybe (a namesake), Беда an
    ordinary criminal case; Иванов's latest news a year old."""
    now = datetime.now(UTC)
    with session_factory.begin() as session:
        snapshot = RosfinmonitoringSnapshotRecord(
            snapshot_date=now,
            source_url="https://example.test",
            content_hash="x",
            entry_count=2,
            fetched_at=now,
        )
        session.add(snapshot)
        session.flush()
        namesake = RosfinmonitoringEntryRecord(
            snapshot_id=snapshot.id,
            full_name="ИВАНОВ ИВАН ИВАНОВИЧ",
            normalized_name="иванов иван иванович",
            matching_key="ивановиваниванович",
        )
        confirmed = RosfinmonitoringEntryRecord(
            snapshot_id=snapshot.id,
            full_name="СМИРНОВА АННА ПЕТРОВНА",
            normalized_name="смирнова анна петровна",
            matching_key="смирновааннапетровна",
        )
        session.add_all([namesake, confirmed])
        session.flush()
        for key, name, verdict, days, entry, level in (
            ("анна смирнова", "Анна Смирнова", "political", 10, confirmed, "full"),
            ("иван иванов", "Иван Иванов", "political", 400, namesake, "name"),
            ("александр беда", "Александр Беда", "criminal", 5, None, None),
        ):
            entity = EntityGroupRecord(
                key=key,
                name=name,
                variants=[[name, 1]],
                mention_count=1,
                article_count=1,
                event_types={},
                regions=[["Москва", 1]] if key == "анна смирнова" else [],
                last_published_at=now - timedelta(days=days),
            )
            session.add(entity)
            session.flush()
            session.add(
                EntityGroupPoliticsRecord(
                    group_id=entity.id,
                    verdict=verdict,
                    method="model",
                    reason=f"так про {name}",
                    quote="цитата",
                )
            )
            if entry is not None and level is not None:
                session.add(
                    EntityGroupRfMatchRecord(group_id=entity.id, entry_id=entry.id, level=level)
                )


def test_the_result_is_the_legacy_list_s(session_factory: sessionmaker[Session]) -> None:
    _political(session_factory)

    with _client(session_factory) as client:
        body = client.get("/api/v1/political").json()
        fresh = client.get("/api/v1/political", params={"months": 3}).json()
        fresh_file = client.get(fresh["export_url"])
        page = client.get("/ui/political", params={"months": 0}).text

    assert body["total"] == 2 and f"Найдено: {body['total']}." in page
    by_name = {row["name"]: row for row in body["items"]}
    assert set(by_name) == {"Смирнова Анна", "Иванов Иван"}
    assert by_name["Смирнова Анна"]["rf_label"] == "в перечне РФМ"
    assert by_name["Смирнова Анна"]["basis"] == "модель: так про Анна Смирнова"
    assert by_name["Смирнова Анна"]["regions"] == "Москва"
    assert by_name["Иванов Иван"]["rf_label"] == "возможно в перечне"
    assert by_name["Смирнова Анна"]["basis"] in page
    # No base loaded: nobody is «not in the base», and the choice is not offered.
    assert body["base_loaded"] is False and body["known"] == []
    assert all(row["known"] is None and not row["not_in_base"] for row in body["items"])
    assert [option["value"] for option in body["periods"]] == ["0", "1", "3", "6", "12"]
    # The latest news a year old is outside three months.
    assert [row["name"] for row in fresh["items"]] == ["Смирнова Анна"]
    # The file the answer names holds the people the answer lists.
    assert fresh["export_url"] == (
        "/ui/political/export.xlsx?months=3&date_from=&date_to=&news=all&known=all&done=hide"
        "&who=all&rfm=all"
    )
    sheet = load_workbook(BytesIO(fresh_file.content))["Результат"]
    assert [row[1] for row in sheet.iter_rows(min_row=2, values_only=True)] == ["Смирнова Анна"]


def test_done_takes_a_person_off_the_result_and_back(
    session_factory: sessionmaker[Session],
) -> None:
    _political(session_factory)

    with _client(session_factory) as client:
        ticked = client.post("/api/v1/political/done", json={"key": "анна смирнова", "done": True})
        hidden = client.get("/api/v1/political").json()
        shown = client.get("/api/v1/political", params={"done": "show"}).json()
        legacy = client.get("/ui/political", params={"months": 0}).text
        unticked = client.post(
            "/api/v1/political/done", json={"key": "анна смирнова", "done": False}
        )
        back = client.get("/api/v1/political").json()
        nobody = client.post("/api/v1/political/done", json={"key": "никто", "done": True})
        foreign = client.post(
            "/api/v1/political/done",
            json={"key": "анна смирнова", "done": True},
            headers={"Origin": "https://evil.example"},
        )
        after_foreign = client.get("/api/v1/political").json()

    assert ticked.status_code == 200 and ticked.json() == {"key": "анна смирнова", "done": True}
    assert [row["name"] for row in hidden["items"]] == ["Иванов Иван"]
    assert hidden["done_total"] == 1
    assert {row["name"]: row["done"] for row in shown["items"]}["Смирнова Анна"] is True
    # The legacy page reads the same mark.
    assert "Показать обработанных (1)" in legacy
    assert unticked.json() == {"key": "анна смирнова", "done": False}
    assert back["total"] == 2 and back["done_total"] == 0
    assert nobody.status_code == 404 and nobody.json() == {"detail": "Человек не найден"}
    # Another site cannot tick for the operator.
    assert foreign.status_code == 403 and after_foreign["done_total"] == 0
