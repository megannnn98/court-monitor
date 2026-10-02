"""«Выгрузить всех людей»: the whole base in one file, whatever a page's filters say."""

from __future__ import annotations

from collections.abc import Iterator
from contextlib import contextmanager
from datetime import UTC, datetime, timedelta
from io import BytesIO
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session, sessionmaker
from support.db_fixtures import DatabaseSeeder

from db.orm_models import (
    AirtableKnownPersonRecord,
    EntityDoneMarkRecord,
    EntityGroupChargeRecord,
    EntityGroupNewsRecord,
    EntityGroupPoliticsRecord,
    EntityGroupRecord,
    EntityGroupRfMatchRecord,
    EntityGroupRoleRecord,
    RosfinmonitoringEntryRecord,
    RosfinmonitoringSnapshotRecord,
    UnnamedFigurantRecord,
    UnnamedIdentityResolutionRecord,
)
from web.app import app
from web.dependencies import get_db

NOW = datetime(2026, 9, 30, 12, 0, tzinfo=UTC)


@contextmanager
def _client(session_factory: sessionmaker[Session]) -> Iterator[TestClient]:
    def override_get_db() -> Iterator[Session]:
        with session_factory() as session:
            yield session

    app.dependency_overrides[get_db] = override_get_db
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.pop(get_db, None)


def _group(session: Session, key: str, name: str, **fields: Any) -> EntityGroupRecord:
    group = EntityGroupRecord(
        key=key,
        name=name,
        variants=[[name, 1]],
        mention_count=fields.pop("mentions", 1),
        article_count=1,
        event_types={},
        regions=fields.pop("regions", []),
        last_published_at=fields.pop("last", NOW),
        **fields,
    )
    session.add(group)
    session.flush()
    return group


def _seed(session_factory: sessionmaker[Session]) -> None:
    with session_factory() as session:
        seed = DatabaseSeeder(session)
        source = seed.source("news", "https://news.example.test")
        article, run = seed.article(
            source, external_id="a1", title="Дело", text="Суд.", published_at=NOW
        )
        event = seed.event(run, "Суд", event_type="arrest", event_date=NOW, links=[])
        snapshot = RosfinmonitoringSnapshotRecord(
            snapshot_date=NOW,
            source_url="https://example.test",
            content_hash="x",
            entry_count=2,
            fetched_at=NOW,
        )
        session.add(snapshot)
        session.flush()
        full = RosfinmonitoringEntryRecord(
            snapshot_id=snapshot.id,
            full_name="СМИРНОВА АННА ПЕТРОВНА",
            normalized_name="смирнова анна петровна",
            matching_key="смирновааннапетровна",
            birth_date=datetime(1990, 2, 1, tzinfo=UTC),
            birth_place="Г. МОСКВА",
            inclusion_date=datetime(2024, 3, 14, tzinfo=UTC),
        )
        namesake = RosfinmonitoringEntryRecord(
            snapshot_id=snapshot.id,
            full_name="ИВАНОВ ИВАН ИВАНОВИЧ",
            normalized_name="иванов иван иванович",
            matching_key="ивановиваниванович",
            inclusion_date=datetime(2020, 1, 1, tzinfo=UTC),
        )
        session.add_all([full, namesake])
        smirnova = _group(
            session,
            "анна смирнова",
            "Анна Смирнова",
            mentions=9,
            regions=[["Москва", 1]],
            last=NOW - timedelta(days=10),
        )
        ivanov = _group(session, "иван иванов", "Иван Иванов", mentions=5)
        judge = _group(session, "петр судьин", "Пётр Судьин", mentions=2)
        # «=SUM(1)»: a scraped name that Excel would read as a formula.
        _group(session, "а-формула", "=SUM(1)", mentions=1)
        for group, verdict in ((smirnova, "political"), (ivanov, "criminal")):
            session.add(
                EntityGroupPoliticsRecord(
                    group_id=group.id, verdict=verdict, method="model", reason="r", quote="q"
                )
            )
        session.add(
            EntityGroupRoleRecord(
                group_id=smirnova.id,
                role="figurant",
                kind=None,
                method="article",
                reason="r",
                quote="q",
            )
        )
        session.add(
            EntityGroupRoleRecord(
                group_id=judge.id,
                role="mentioned",
                kind="judge",
                method="model",
                reason="r",
                quote="q",
            )
        )
        session.add(
            EntityGroupNewsRecord(
                group_id=smirnova.id,
                kind="new_case",
                method="model",
                reason="возбуждено",
                quote="q",
                published_at=NOW,
            )
        )
        for article_number in ("205.2", "282", "30", "63", "205.2"):
            session.add(
                EntityGroupChargeRecord(
                    group_id=smirnova.id,
                    event_id=event,
                    publication_id=article,
                    article=article_number,
                    event_type="arrest",
                    other_targets=0,
                    quote="q",
                )
            )
        undated = RosfinmonitoringEntryRecord(
            snapshot_id=snapshot.id,
            full_name="СУДЬИН ПЁТР ПЕТРОВИЧ",
            normalized_name="судьин петр петрович",
            matching_key="судьинпетрпетрович",
        )
        session.add(undated)
        session.flush()
        session.add(EntityGroupRfMatchRecord(group_id=judge.id, entry_id=undated.id, level="full"))
        session.add(EntityGroupRfMatchRecord(group_id=smirnova.id, entry_id=full.id, level="full"))
        session.add(
            EntityGroupRfMatchRecord(group_id=ivanov.id, entry_id=namesake.id, level="name")
        )
        session.add(
            AirtableKnownPersonRecord(
                external_id="k1",
                full_name="Смирнова Анна Петровна",
                normalized_name="смирнова анна петровна",
                matching_key="смирновааннапетровна",
                active=True,
            )
        )
        # Done, and still hidden on the result: the file holds it all the same.
        session.add(EntityDoneMarkRecord(key="иван иванов", news_at=NOW))
        session.add(
            EntityDoneMarkRecord(key="unnamed:канаш|male|15|", news_at=NOW - timedelta(days=5))
        )
        for key, age, place, articles, day in (
            ("u1", 15, "Канаш", ["205.5"], 5),
            ("u2", 30, "Тотьма", [], 6),
            ("u3", 40, "Керчь", ["275"], 7),
        ):
            quote = f"Задержан человек {key}."
            doc, _ = seed.article(
                source,
                external_id=key,
                title=f"Новость {key}",
                text=quote,
                published_at=NOW - timedelta(days=day),
            )
            session.add(
                UnnamedFigurantRecord(
                    key=key,
                    article_id=doc,
                    start_offset=0,
                    end_offset=len(quote),
                    quote=quote,
                    age=age,
                    gender="male",
                    place=place,
                    initial=None,
                    articles=articles,
                    event_type="detention",
                    explanation=f"объяснение {key}",
                    published_at=NOW - timedelta(days=day),
                )
            )
        session.add(UnnamedIdentityResolutionRecord(figurant_key="u3", resolution="supplied_name"))
        session.commit()


def _book(session_factory: sessionmaker[Session]) -> Any:
    with _client(session_factory) as client:
        response = client.get("/ui/people/export.xlsx")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert response.headers["content-disposition"].startswith('attachment; filename="people_')
    return load_workbook(BytesIO(response.content))


def test_the_file_holds_every_person_with_a_name_the_most_mentioned_first(
    session_factory: sessionmaker[Session],
) -> None:
    _seed(session_factory)
    book = _book(session_factory)

    assert book.sheetnames == ["Люди", "Без имени", "Пояснения"]
    rows = list(book["Люди"].iter_rows(values_only=True))
    assert [cell for cell in rows[0]][:5] == [
        "№",
        "Фамилия Имя",
        "Роль в деле",
        "Дело",
        "Статьи УК",
    ]
    assert [row[1] for row in rows[1:]] == [
        "Смирнова Анна",
        "Иванов Иван",
        "Судьин Пётр",
        "=SUM(1)",
    ]
    smirnova, ivanov, judge, formula = rows[1:]

    # In the order of the numbers, not of the letters: 30 and 63 before 205.2.
    assert smirnova[2:7] == ("фигурант дела", "политическое", "30, 63, 205.2, 282", "Москва", 9)
    assert smirnova[8].date() == (NOW - timedelta(days=10)).date()
    assert smirnova[9] == "новое дело"
    assert smirnova[10].startswith("в перечне: СМИРНОВА АННА ПЕТРОВНА, 01.02.1990 г.р., Г. МОСКВА")
    assert smirnova[10].endswith("запись перечня включена 14.03.2024")
    # The news give two words and the base three: one record, but not certainly the person.
    assert smirnova[11] == "вероятно, есть в базе: Смирнова Анна Петровна"
    # A namesake: the entry's day is somebody else's and is not in the file.
    assert ivanov[3] == "обычное уголовное" and ivanov[10] == "возможно тёзка: ИВАНОВ ИВАН ИВАНОВИЧ"
    assert "2020" not in ivanov[10]
    assert ivanov[12] == "да" and smirnova[12] is None
    assert judge[2] == "упомянут: судья" and judge[3] is None, "never judged: an empty cell"
    assert ivanov[2] is None and formula[2] is None, "no role yet: an empty cell"
    # An entry with no day: the words say no day, and no dangling comma.
    assert judge[10] == "в перечне: СУДЬИН ПЁТР ПЕТРОВИЧ"
    assert formula[1] == "=SUM(1)" and book["Люди"].cell(row=5, column=2).data_type == "s"


def test_the_file_holds_what_the_pages_hide(session_factory: sessionmaker[Session]) -> None:
    """Иванов is «обработано» and «уголовное»: not on the result, and in the file."""
    _seed(session_factory)
    with _client(session_factory) as client:
        page = client.get("/ui/political?months=0&who=named").text
    assert "Иванов Иван" not in page

    names = [
        row[1] for row in _book(session_factory)["Люди"].iter_rows(min_row=2, values_only=True)
    ]
    assert "Иванов Иван" in names


def test_the_sheet_of_the_unnamed_holds_every_open_case_political_or_not(
    session_factory: sessionmaker[Session],
) -> None:
    _seed(session_factory)
    rows = list(_book(session_factory)["Без имени"].iter_rows(values_only=True))

    assert [cell for cell in rows[0]][:4] == ["№", "Кто", "Политическая статья", "Статьи УК"]
    # The latest first; the identified one (u3) is a person with a name now.
    assert [row[1] for row in rows[1:]] == ["15-летний житель Канаша", "30-летний житель Тотьмы"]
    kanash, totma = rows[1:]
    assert kanash[2:6] == ("да", "205.5", "Канаш", 1) and totma[2:4] == ("нет", None)
    assert kanash[11] == "да" and totma[11] is None, "done on the site, done in the file"
    assert kanash[8] == "задержание" and kanash[9] == "объяснение u1"
    assert kanash[10] == "news: Новость u1" and kanash[7].date() == (NOW - timedelta(days=5)).date()


def test_the_source_of_an_unnamed_case_is_a_link_and_the_dates_are_dates(
    session_factory: sessionmaker[Session],
) -> None:
    _seed(session_factory)
    sheet = _book(session_factory)["Без имени"]

    assert sheet.cell(row=2, column=11).hyperlink is not None
    assert sheet.cell(row=2, column=8).number_format == "DD.MM.YYYY"
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref is not None


def test_the_notes_say_what_the_file_is_and_carry_the_attribution(
    session_factory: sessionmaker[Session],
) -> None:
    _seed(session_factory)
    notes = " ".join(
        str(row[0]) for row in _book(session_factory)["Пояснения"].iter_rows(values_only=True)
    )

    assert "без фильтров" in notes and "ОВД-Инфо (CC BY 3.0)" in notes
    assert "не делает совпадение подтверждением личности" in notes


def test_an_empty_base_gives_a_file_with_headers_only(
    session_factory: sessionmaker[Session],
) -> None:
    book = _book(session_factory)

    assert [row[0] for row in book["Люди"].iter_rows(values_only=True)] == ["№"]
    assert [row[0] for row in book["Без имени"].iter_rows(values_only=True)] == ["№"]


def test_the_people_page_has_the_button(session_factory: sessionmaker[Session]) -> None:
    with _client(session_factory) as client:
        page = client.get("/ui/entities").text

    assert 'href="/ui/people/export.xlsx"' in page and "Выгрузить всех" in page
