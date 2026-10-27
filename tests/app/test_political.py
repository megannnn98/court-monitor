"""«Результат»: the politically persecuted off the Rosfinmonitoring list, and its Excel.

The list is read and ticked through `/api/v1/political`, as the React page does; the
Excel file is the server's, at the address the API names."""

from __future__ import annotations

from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import replace
from datetime import UTC, date, datetime, timedelta
from io import BytesIO
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session, sessionmaker
from support.db_fixtures import DatabaseSeeder

from analysis.read import Politics
from analysis.storage.models import (
    AirtableKnownPersonRecord,
    EntityGroupNewsRecord,
    EntityGroupPoliticsRecord,
    EntityGroupRfMatchRecord,
    RosfinmonitoringEntryRecord,
    RosfinmonitoringSnapshotRecord,
    UnnamedFigurantRecord,
    UnnamedIdentityResolutionRecord,
)
from processing.read import Group
from processing.storage.models import EntityGroupRecord
from results.storage.models import EntityDoneMarkRecord
from web.app import app
from web.dependencies import get_db
from web.views.political import political_xlsx
from web.views.political_rows import named_row

API = "/api/v1/political"
EXPORT = "/ui/political/export.xlsx"


@contextmanager
def _client(session_factory: sessionmaker[Session]) -> Iterator[TestClient]:
    def override_get_db() -> Iterator[Session]:
        with session_factory() as session:
            yield session

    app.dependency_overrides[get_db] = override_get_db
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.pop(get_db, None)


def _list(client: TestClient, **params: str | int) -> dict[str, Any]:
    listed: dict[str, Any] = client.get(API, params=params).json()
    return listed


def _names(listed: dict[str, Any]) -> list[str]:
    return [row["name"] for row in listed["items"]]


def _row(listed: dict[str, Any], name: str) -> dict[str, Any]:
    row: dict[str, Any] = next(item for item in listed["items"] if name in item["name"])
    return row


def _counts(options: list[dict[str, Any]]) -> dict[str, int | None]:
    return {option["value"]: option["count"] for option in options}


def _seed(session_factory: sessionmaker[Session]) -> None:
    now = datetime.now(UTC)
    with session_factory.begin() as session:
        snapshot = RosfinmonitoringSnapshotRecord(
            snapshot_date=now,
            source_url="https://example.test",
            content_hash="x",
            entry_count=1,
            fetched_at=now,
        )
        session.add(snapshot)
        session.flush()
        entry = RosfinmonitoringEntryRecord(
            snapshot_id=snapshot.id,
            full_name="ИВАНОВ ИВАН ИВАНОВИЧ",
            normalized_name="иванов иван иванович",
            matching_key="ивановиваниванович",
        )
        session.add(entry)
        # Смирнова on the list with her patronymic: who she is, confirmed.
        confirmed = RosfinmonitoringEntryRecord(
            snapshot_id=snapshot.id,
            full_name="СМИРНОВА АННА ПЕТРОВНА",
            normalized_name="смирнова анна петровна",
            matching_key="смирновааннапетровна",
            birth_date=datetime(1990, 2, 1, tzinfo=UTC),
            birth_place="Г. МОСКВА",
        )
        session.add(confirmed)
        for key, name, verdict, days, regions in (
            ("анна смирнова", "Анна Смирнова", "political", 10, [["Москва", 1]]),
            ("иван иванов", "Иван Иванов", "political", 400, []),
            ("александр беда", "Александр Беда", "criminal", 5, []),
        ):
            entity = EntityGroupRecord(
                key=key,
                name=name,
                variants=[[name, 1]],
                mention_count=1,
                article_count=1,
                event_types={},
                regions=regions,
                last_published_at=now - timedelta(days=days),
            )
            session.add(entity)
            session.flush()
            session.add(
                EntityGroupPoliticsRecord(
                    group_id=entity.id,
                    verdict=verdict,
                    method="model",
                    reason=f"так про {name}",
                    quote="цитата",
                )
            )
            if key == "иван иванов":
                session.flush()
                session.add(
                    EntityGroupRfMatchRecord(group_id=entity.id, entry_id=entry.id, level="name")
                )
            if key == "анна смирнова":
                session.flush()
                session.add(
                    EntityGroupRfMatchRecord(
                        group_id=entity.id, entry_id=confirmed.id, level="full"
                    )
                )


def test_the_list_is_the_political_the_period_tells_new_from_old(
    session_factory: sessionmaker[Session],
) -> None:
    _seed(session_factory)

    with _client(session_factory) as client:
        everything = _list(client)
        fresh = _list(client, months=3)
        # The box that hid the maybe-namesakes is gone: an old address hides nobody.
        old_box = _list(client, months=0, hide_maybe_listed="true")

    assert everything["total"] == 2 and _names(everything) == ["Смирнова Анна", "Иванов Иван"]
    smirnova = _row(everything, "Смирнова")
    assert smirnova["regions"] == "Москва" and smirnova["basis"] == "модель: так про Анна Смирнова"
    # On the list with the patronymic: in the result, the list's word beside the name.
    assert smirnova["rf_label"] == "в перечне РФМ"
    assert smirnova["rf_entry"] == "СМИРНОВА АННА ПЕТРОВНА, 01.02.1990 г.р., Г. МОСКВА"
    assert _row(everything, "Иванов")["rf_label"] == "возможно в перечне"
    # Иванов's latest news is a year old: not a new case.
    assert fresh["total"] == 1 and _names(fresh) == ["Смирнова Анна"]
    assert old_box["total"] == 2 and "Иванов Иван" in _names(old_box)


def test_the_excel_has_every_row_of_the_filters(session_factory: sessionmaker[Session]) -> None:
    _seed(session_factory)

    with _client(session_factory) as client:
        response = client.get(EXPORT, params={"months": 0})

    assert response.status_code == 200
    # All the time: from the earliest latest news (Иванов's, 400 days ago) to today.
    today = datetime.now(UTC).date()
    assert response.headers["content-disposition"] == (
        f'attachment; filename="result_{(today - timedelta(days=400)).isoformat()}_'
        f'{today.isoformat()}.xlsx"'
    )
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:3] == ("№", "Фамилия Имя", "Регион")
    assert rows[0][8] == "Перечень РФМ"
    assert rows[0][9:14] == (
        "Источник 1",
        "Источник 2",
        "Источник 3",
        "Свежая новость",
        "В базе Airtable",
    )
    assert [(row[1], row[2], row[8]) for row in rows[1:]] == [
        (
            "Смирнова Анна",
            "Москва",
            "в перечне: СМИРНОВА АННА ПЕТРОВНА, 01.02.1990 г.р., Г. МОСКВА",
        ),
        ("Иванов Иван", None, "возможно тёзка: ИВАНОВ ИВАН ИВАНОВИЧ"),
    ]


def test_the_api_names_the_excel_of_its_own_filters(
    session_factory: sessionmaker[Session],
) -> None:
    _seed(session_factory)

    with _client(session_factory) as client:
        listed = _list(client, months=3, news="sentence")
        response = client.get(listed["export_url"])

    assert listed["export_url"].startswith(f"{EXPORT}?") and "news=sentence" in listed["export_url"]
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet is not None and list(sheet.iter_rows(min_row=2, values_only=True)) == []


def test_the_excel_makes_each_source_a_separate_link() -> None:
    entity = Group(
        id=1,
        key="анна смирнова",
        name="Анна Смирнова",
        mention_count=1,
        article_count=1,
        last_published_at=None,
        gender=None,
        name_source="rules",
        collected_at=datetime(2026, 10, 1, tzinfo=UTC),
    )
    politics = Politics(
        group_id=1,
        verdict="political",
        method="model",
        reason="основание",
        quote="цитата",
    )
    workbook = load_workbook(
        BytesIO(
            political_xlsx(
                [
                    replace(
                        named_row(entity, politics),
                        links=[
                            ("Первый источник", "https://first.example.test", None, "Первый"),
                            ("Второй источник", "https://second.example.test", None, "Второй"),
                        ],
                    )
                ]
            )
        )
    )
    sheet = workbook.active
    assert sheet is not None
    assert sheet.cell(row=2, column=10).value == "Первый: Первый источник"
    first_link = sheet.cell(row=2, column=10).hyperlink
    assert first_link is not None
    assert first_link.target == "https://first.example.test"
    assert sheet.cell(row=2, column=11).value == "Второй: Второй источник"
    second_link = sheet.cell(row=2, column=11).hyperlink
    assert second_link is not None
    assert second_link.target == "https://second.example.test"
    assert sheet.cell(row=2, column=12).value is None


def test_a_period_of_dates(session_factory: sessionmaker[Session]) -> None:
    _seed(session_factory)
    today = datetime.now(UTC).date()
    recent = {"date_from": (today - timedelta(days=30)).isoformat(), "date_to": today.isoformat()}
    old = {
        "date_from": (today - timedelta(days=500)).isoformat(),
        "date_to": (today - timedelta(days=300)).isoformat(),
    }

    with _client(session_factory) as client:
        in_recent = _list(client, **recent)
        in_old = _list(client, **old)
        # Dates given, the months are not the choice: a year-old case, though «3 months».
        dates_win = _list(client, **old, months=3)
        excel = client.get(EXPORT, params=old)

    # Смирнова's latest news is 10 days old, Иванов's 400.
    assert in_recent["total"] == 1 and _names(in_recent) == ["Смирнова Анна"]
    assert in_old["total"] == 1 and _names(in_old) == ["Иванов Иван"]
    assert _names(dates_win) == ["Иванов Иван"]
    assert (
        f'filename="result_{old["date_from"]}_{old["date_to"]}.xlsx"'
        in (excel.headers["content-disposition"])
    )
    sheet = load_workbook(BytesIO(excel.content)).active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))
    assert [row[1] for row in rows[1:]] == ["Иванов Иван"]


def test_the_form_s_day_month_year_is_read(session_factory: sessionmaker[Session]) -> None:
    _seed(session_factory)
    today = datetime.now(UTC).date()
    written = {
        "date_from": (today - timedelta(days=30)).strftime("%d/%m/%Y"),
        "date_to": today.strftime("%d/%m/%Y"),
    }

    with _client(session_factory) as client:
        listed = _list(client, **written)

    # Смирнова's latest news is 10 days old, Иванов's 400.
    assert listed["total"] == 1 and _names(listed) == ["Смирнова Анна"]


def _news(session_factory: sessionmaker[Session]) -> None:
    """Смирнова's latest news: a new case. Иванов's: never read."""
    with session_factory.begin() as session:
        smirnova = session.query(EntityGroupRecord).filter_by(key="анна смирнова").one()
        session.add(
            EntityGroupNewsRecord(
                group_id=smirnova.id,
                kind="new_case",
                method="model",
                reason="возбуждено дело",
                quote="цитата",
                published_at=smirnova.last_published_at,
            )
        )


def test_the_latest_news_is_marked_and_chosen(session_factory: sessionmaker[Session]) -> None:
    _seed(session_factory)
    _news(session_factory)

    with _client(session_factory) as client:
        everything = _list(client, months=0)
        new_cases = _list(client, months=0, news="new_case")
        sentences = _list(client, months=0, news="sentence")
        nonsense = _list(client, months=0, news="drop table")
        excel = client.get(EXPORT, params={"months": 0, "news": "new_case"})

    smirnova = _row(everything, "Смирнова")
    assert (smirnova["news_kind"], smirnova["news_label"], smirnova["news_reason"]) == (
        "new_case",
        "новое дело",
        "возбуждено дело",
    )
    assert _row(everything, "Иванов")["news_kind"] is None
    news = _counts(everything["news"])
    assert news["new_case"] == 1 and news["sentence"] == 0
    assert new_cases["total"] == 1 and _names(new_cases) == ["Смирнова Анна"]
    assert sentences["total"] == 0
    assert nonsense["total"] == 2
    sheet = load_workbook(BytesIO(excel.content)).active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))
    column = rows[0].index("Свежая новость")
    assert [(row[1], row[column]) for row in rows[1:]] == [("Смирнова Анна", "новое дело")]


def _known(session_factory: sessionmaker[Session], *names: str, active: bool = True) -> None:
    with session_factory.begin() as session:
        for number, name in enumerate(names):
            session.add(
                AirtableKnownPersonRecord(
                    external_id=f"{name}-{number}",
                    full_name=name,
                    normalized_name=name.lower(),
                    matching_key=name.lower().replace(" ", ""),
                    active=active,
                )
            )


def _third(session_factory: sessionmaker[Session]) -> None:
    """A political person nobody has: the base's answer for them is «not there»."""
    with session_factory.begin() as session:
        entity = EntityGroupRecord(
            key="олег новиков",
            name="Олег Новиков",
            variants=[["Олег Новиков", 1]],
            mention_count=1,
            article_count=1,
            event_types={},
            regions=[],
            last_published_at=datetime.now(UTC) - timedelta(days=2),
        )
        session.add(entity)
        session.flush()
        session.add(
            EntityGroupPoliticsRecord(
                group_id=entity.id, verdict="political", method="model", reason="", quote=""
            )
        )


def _base_world(session_factory: sessionmaker[Session]) -> None:
    """Смирнова: one record, with a patronymic the news lacks. Иванов: two namesakes.
    Новиков: not in the base."""
    _seed(session_factory)
    _third(session_factory)
    _known(
        session_factory,
        "Смирнова Анна Петровна",
        "Иванов Иван Иванович",
        "Иванов Иван Петрович",
        "Кто-то Другой",
    )


def test_the_list_says_who_the_operator_s_base_already_holds(
    session_factory: sessionmaker[Session],
) -> None:
    _base_world(session_factory)

    with _client(session_factory) as client:
        listed = _list(client, months=0)

    assert listed["base_loaded"] is True
    # Смирнова: one record that holds her name and more — probably, not certainly.
    smirnova = _row(listed, "Смирнова")
    assert smirnova["known"] == {
        "level": "probably",
        "label": "вероятно, есть в базе",
        "names": ["Смирнова Анна Петровна"],
    }
    # Иванов: two records fit; the list counts them and picks neither.
    ivanov = _row(listed, "Иванов")
    assert ivanov["known"]["level"] == "namesakes"
    assert ivanov["known"]["label"] == "тёзки в базе: 2"
    # Новиков: nobody by this name — the new one.
    novikov = _row(listed, "Новиков")
    assert novikov["known"] is None and novikov["not_in_base"] is True
    known = _counts(listed["known"])
    assert (known["none"], known["probably"], known["namesakes"], known["in_base"]) == (1, 1, 1, 0)


def test_the_choice_of_an_answer_narrows_the_list_and_its_count(
    session_factory: sessionmaker[Session],
) -> None:
    _base_world(session_factory)

    with _client(session_factory) as client:
        new = _list(client, months=0, known="none")
        namesakes = _list(client, months=0, known="namesakes")
        nonsense = _list(client, months=0, known="drop table")

    assert new["total"] == 1 and _names(new) == ["Новиков Олег"]
    assert namesakes["total"] == 1 and _names(namesakes) == ["Иванов Иван"]
    # Narrowed, the counts stay those of the whole list: the choice is not a filter on them.
    assert _counts(new["known"])["probably"] == 1 and _counts(new["known"])["none"] == 1
    assert nonsense["total"] == 3


def test_with_no_base_loaded_nobody_is_called_new(session_factory: sessionmaker[Session]) -> None:
    """Nothing synced yet: «not in the base» would be a statement about nothing."""
    _seed(session_factory)

    with _client(session_factory) as client:
        # An answer that, taken at its word, would leave nobody: with no base it is ignored.
        listed = _list(client, months=0, known="probably")
        excel = client.get(EXPORT, params={"months": 0, "known": "probably"})

    assert listed["base_loaded"] is False and listed["known"] == []
    assert listed["total"] == 2
    assert all(row["known"] is None and not row["not_in_base"] for row in listed["items"])
    sheet = load_workbook(BytesIO(excel.content)).active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))
    assert len(rows) == 3 and all(row[-1] is None for row in rows[1:])


def test_the_excel_has_the_base_s_answer_and_follows_the_choice(
    session_factory: sessionmaker[Session],
) -> None:
    _base_world(session_factory)

    with _client(session_factory) as client:
        everyone = client.get(EXPORT, params={"months": 0})
        new = client.get(EXPORT, params={"months": 0, "known": "none"})

    def answers(response: object) -> dict[str, str | None]:
        content = getattr(response, "content", None)
        assert isinstance(content, bytes)
        sheet = load_workbook(BytesIO(content)).active
        assert sheet is not None
        rows = list(sheet.iter_rows(values_only=True))
        column = rows[0].index("В базе Airtable")
        answer: dict[str, str | None] = {}
        for row in rows[1:]:
            assert isinstance(row[1], str)
            value = row[column]
            assert value is None or isinstance(value, str)
            answer[row[1]] = value
        return answer

    assert answers(everyone) == {
        "Смирнова Анна": "вероятно, есть в базе: Смирнова Анна Петровна",
        "Иванов Иван": "тёзки в базе: 2 (Иванов Иван Иванович; Иванов Иван Петрович)",
        "Новиков Олег": "нет в базе",
    }
    assert answers(new) == {"Новиков Олег": "нет в базе"}


def test_an_inactive_record_of_the_base_does_not_count(
    session_factory: sessionmaker[Session],
) -> None:
    _seed(session_factory)
    _known(session_factory, "Смирнова Анна Петровна", active=False)
    _known(session_factory, "Кто-то Другой")

    with _client(session_factory) as client:
        listed = _list(client, months=0)

    assert listed["base_loaded"] is True
    smirnova = _row(listed, "Смирнова")
    assert smirnova["known"] is None and smirnova["not_in_base"] is True


def _tick(client: TestClient, key: str, done: bool = True) -> None:
    response = client.post(f"{API}/done", json={"key": key, "done": done})
    assert response.status_code == 200, response.text
    assert response.json() == {"key": key, "done": done}


def test_a_person_ticked_as_done_leaves_the_list_and_can_be_shown_again(
    session_factory: sessionmaker[Session],
) -> None:
    """Ирина ticks the people she has dealt with, so the list holds only what is left."""
    _seed(session_factory)
    with _client(session_factory) as client:
        before = _list(client)
        assert _names(before) == ["Смирнова Анна", "Иванов Иван"]
        assert before["done_total"] == 0, "the way back is always in sight"

        _tick(client, "анна смирнова")
        hidden = _list(client, months=0)
        assert _names(hidden) == ["Иванов Иван"] and hidden["done_total"] == 1

        shown = _list(client, months=0, done="show")
        assert shown["total"] == 2 and _row(shown, "Смирнова")["done"] is True
        assert _row(shown, "Иванов")["done"] is False

        # The file follows the list: hidden there, hidden here.
        book = load_workbook(BytesIO(client.get(f"{EXPORT}?months=0").content))
        names = [row[1] for row in book["Результат"].iter_rows(min_row=2, values_only=True)]
        assert names == ["Иванов Иван"]

        _tick(client, "анна смирнова", done=False)
        assert "Смирнова Анна" in _names(_list(client, months=0))


def test_a_later_news_brings_a_done_person_back(session_factory: sessionmaker[Session]) -> None:
    """What was dealt with is what was known then. A news after the tick is new work."""
    _seed(session_factory)
    with _client(session_factory) as client:
        _tick(client, "анна смирнова")
        assert "Смирнова Анна" not in _names(_list(client, months=0))
        with session_factory.begin() as session:
            entity = session.query(EntityGroupRecord).filter_by(key="анна смирнова").one()
            entity.last_published_at = datetime.now(UTC)
        assert "Смирнова Анна" in _names(_list(client, months=0))


def test_the_mark_survives_a_rebuild_of_the_people(session_factory: sessionmaker[Session]) -> None:
    """Step 3 rebuilds the groups and their ids change: the mark is kept by the key."""
    _seed(session_factory)
    with _client(session_factory) as client:
        _tick(client, "анна смирнова")
        with session_factory.begin() as session:
            marks = session.query(EntityDoneMarkRecord).all()
            assert [(mark.key, mark.news_at is not None) for mark in marks] == [
                ("анна смирнова", True)
            ]


def test_ticking_nobody_is_refused(session_factory: sessionmaker[Session]) -> None:
    _seed(session_factory)
    with _client(session_factory) as client:
        missing = client.post(f"{API}/done", json={"key": "никто", "done": True})
        foreign = client.post(
            f"{API}/done",
            json={"key": "иван иванов", "done": True},
            headers={"Origin": "https://evil.example"},
        )
        listed = _list(client, months=0)

    assert missing.status_code == 404 and missing.json()["detail"] == "Человек не найден"
    assert foreign.status_code == 403 and listed["done_total"] == 0


def _unnamed(
    session_factory: sessionmaker[Session],
    key: str,
    *,
    age: int | None,
    place: str,
    articles: list[str],
    gender: str = "male",
    initial: str | None = None,
    event: str = "detention",
    day: int = 20,
) -> None:
    quote = f"Задержан человек {key}."
    with session_factory() as session:
        seed = DatabaseSeeder(session)
        source = seed.source(f"news-{key}", f"https://{key}.example.test")
        article, _ = seed.article(
            source,
            external_id=key,
            title=f"Новость {key}",
            text=quote,
            published_at=datetime(2026, 9, day, tzinfo=UTC),
        )
        session.add(
            UnnamedFigurantRecord(
                key=key,
                article_id=article,
                start_offset=0,
                end_offset=len(quote),
                quote=quote,
                age=age,
                gender=gender,
                place=place,
                initial=initial,
                articles=articles,
                event_type=event,
                explanation=f"дело об участии в запрещённой организации ({key})",
                published_at=datetime(2026, 9, day, tzinfo=UTC),
            )
        )
        session.commit()


def test_the_unnamed_with_a_political_article_are_rows_of_the_list_and_of_the_file(
    session_factory: sessionmaker[Session],
) -> None:
    """Ирина: «в работе они есть, а в результатах нет». One person a row, named by what
    the text tells; a case with no political article stays on «Безымянные»."""
    _seed(session_factory)
    _unnamed(session_factory, "a", age=15, place="Канаш", articles=["205.5"], day=20)
    _unnamed(session_factory, "b", age=15, place="Канаш", articles=[], day=21)
    _unnamed(session_factory, "c", age=None, place="Благовещенск", articles=["282.2"])
    _unnamed(session_factory, "d", age=16, place="Вологда", articles=["228.1"])
    _unnamed(session_factory, "e", age=30, place="Тотьма", articles=[])
    with _client(session_factory) as client:
        listed = _list(client, months=0)
        assert listed["total"] == 4
        unnamed = [row for row in listed["items"] if row["unnamed"]]
        assert sorted(row["name"] for row in unnamed) == [
            "15-летний житель Канаша",
            "житель Благовещенска",
        ], "two sentences, one row; Вологда and Тотьма are no political case"
        kanash = _row(listed, "Канаша")
        assert kanash["url"] == "/ui/unnamed?status=all#u-b", "the latest sentence"

        only = _list(client, months=0, who="unnamed")
        assert only["total"] == 2 and all(row["unnamed"] for row in only["items"])
        named = _list(client, months=0, who="named")
        assert named["total"] == 2 and not any(row["unnamed"] for row in named["items"])

        book = load_workbook(BytesIO(client.get(f"{EXPORT}?months=0&who=unnamed").content))
        rows = list(book["Результат"].iter_rows(min_row=2, values_only=True))
        labels = [label for row in rows if isinstance(label := row[1], str)]
        assert len(labels) == len(rows)
        assert sorted(labels) == ["15-летний житель Канаша", "житель Благовещенска"]
        kanash_row = next(row for row in rows if row[1] == "15-летний житель Канаша")
        assert kanash_row[2] == "Канаш" and kanash_row[3] == "205.5"
        assert isinstance(kanash_row[6], datetime) and isinstance(kanash_row[7], datetime)
        assert (kanash_row[6].date(), kanash_row[7].date()) == (
            date(2026, 9, 20),
            date(2026, 9, 21),
        )
        assert isinstance(kanash_row[9], str) and isinstance(kanash_row[10], str)
        assert kanash_row[9].endswith("Новость b") and kanash_row[10].endswith("Новость a")


def test_an_identified_figurant_is_no_longer_an_unnamed_row(
    session_factory: sessionmaker[Session],
) -> None:
    """Once the operator says who it is, the person is in the result under the name."""
    _seed(session_factory)
    _unnamed(session_factory, "a", age=15, place="Канаш", articles=["205.5"])
    _unnamed(session_factory, "c", age=None, place="Благовещенск", articles=["282.2"])
    with session_factory.begin() as session:
        session.add(
            UnnamedIdentityResolutionRecord(
                figurant_key="a", resolution="supplied_name", normalized_name="Иван Петров"
            )
        )
        session.add(UnnamedIdentityResolutionRecord(figurant_key="c", resolution="no_rf_match"))
    with _client(session_factory) as client:
        listed = _list(client, months=0, who="unnamed")
    # «нет записи РФМ» does not close the case.
    assert _names(listed) == ["житель Благовещенска"]


def test_an_unnamed_row_is_ticked_done_and_comes_back_with_a_later_sentence(
    session_factory: sessionmaker[Session],
) -> None:
    _seed(session_factory)
    _unnamed(session_factory, "a", age=15, place="Канаш", articles=["205.5"], day=20)
    key = "unnamed:канаш|male|15|"
    with _client(session_factory) as client:
        assert _row(_list(client, months=0), "Канаша")["key"] == key
        _tick(client, key)
        hidden = _list(client, months=0, who="unnamed")
        assert hidden["items"] == [] and hidden["done_total"] == 1
        shown = _list(client, months=0, who="unnamed", done="show")
        assert _row(shown, "Канаша")["done"] is True

        _unnamed(session_factory, "b", age=15, place="Канаш", articles=[], day=25)
        assert _names(_list(client, months=0, who="unnamed")) == ["15-летний житель Канаша"]

        missing = client.post(f"{API}/done", json={"key": "unnamed:нигде|male|1|", "done": True})
        assert missing.status_code == 404


def test_the_filters_of_the_list_hold_for_the_unnamed_too(
    session_factory: sessionmaker[Session],
) -> None:
    _seed(session_factory)
    _unnamed(session_factory, "a", age=15, place="Канаш", articles=["205.5"], event="sentence")
    _unnamed(session_factory, "c", age=None, place="Благовещенск", articles=["282.2"], day=2)
    with _client(session_factory) as client:
        sentences = _list(client, months=0, who="unnamed", news="sentence")
        assert _names(sentences) == ["15-летний житель Канаша"]
        late = _list(client, who="unnamed", date_from="10/09/2026")
        assert _names(late) == ["15-летний житель Канаша"]
        # With no record of the base that tells the same, an unnamed person is new too.
        _known(session_factory, "Петров Пётр Петрович")
        new = _list(client, months=0, known="none")
        assert {"15-летний житель Канаша", "житель Благовещенска"} <= set(_names(new))


def _nameless(
    session_factory: sessionmaker[Session], name: str, *, opened: date, **values: object
) -> None:
    """A record of the base that names nobody: no birth date, the age in its name."""
    with session_factory.begin() as session:
        session.add(
            AirtableKnownPersonRecord(
                external_id=f"rec-{name}-{opened}",
                full_name=name,
                normalized_name=name.lower(),
                matching_key=name.lower(),
                case_opened_on=opened,
                **{"gender": "male", "region": "Чувашия", "city": "Канаш", **values},
            )
        )


def test_an_unnamed_person_is_looked_for_among_the_base_s_nameless_records(
    session_factory: sessionmaker[Session],
) -> None:
    """«41-летний житель Ленинградской области» stands in the base under that very name;
    the result said «нет в базе», having no name to look up."""
    _seed(session_factory)
    _unnamed(session_factory, "a", age=41, place="Ленинградская область", articles=["205.5"])
    _unnamed(session_factory, "b", age=15, place="Канаш", articles=["205.5"])
    _unnamed(session_factory, "c", age=30, place="Вологда", articles=["282.2"])
    _unnamed(session_factory, "d", age=None, place="Тотьма", articles=["282.2"])
    # The base is loaded: it holds named people too.
    _known(session_factory, "Петров Пётр Петрович")
    _nameless(
        session_factory,
        "41-летний житель <Ленинградской> области",
        opened=date(2026, 9, 19),
        region="Ленинградская область",
        city=None,
    )
    # Two records tell of a boy of Канаш: which one, if any, is for the operator.
    _nameless(session_factory, "15-летний житель Канаша", opened=date(2026, 9, 1))
    _nameless(session_factory, "Подросток 15-летний, Канаш (2)", opened=date(2026, 8, 1))
    # A record the base no longer counts says nothing.
    _nameless(
        session_factory,
        "30-летний житель Вологды",
        opened=date(2026, 9, 1),
        region="Вологодская область",
        city="Вологда",
        active=False,
    )
    # Not them: a woman, another town, a named person of that age and place.
    _nameless(
        session_factory, "15-летняя жительница Канаша", opened=date(2026, 9, 1), gender="female"
    )
    _nameless(
        session_factory,
        "30-летний житель Омска",
        opened=date(2026, 9, 1),
        region="Омская область",
        city="Омск",
    )
    with session_factory.begin() as session:
        session.add(
            AirtableKnownPersonRecord(
                external_id="named",
                full_name="Иванов 15-летний Иван",
                normalized_name="x",
                matching_key="x",
                gender="male",
                region="Чувашия",
                city="Канаш",
                birth_date=date(2011, 1, 1),
            )
        )

    with _client(session_factory) as client:
        listed = _list(client, months=0, who="unnamed")
        similar = _list(client, months=0, known="similar")
        workbook = load_workbook(BytesIO(client.get(f"{EXPORT}?months=0&who=unnamed").content))

    # One record fits: probably her record, named as the base writes it.
    assert _row(listed, "Ленинградской области")["known"] == {
        "level": "probably",
        "label": "вероятно, есть в базе",
        "names": ["41-летний житель <Ленинградской> области"],
    }
    assert _row(listed, "Канаша")["known"] == {
        "level": "similar",
        "label": "похожие записи в базе: 2",
        "names": ["15-летний житель Канаша", "Подросток 15-летний, Канаш (2)"],
    }
    # Nobody of Вологда in the base; and no age told — nothing to compare.
    for place in ("Вологды", "Тотьмы"):
        assert _row(listed, place)["known"] is None and _row(listed, place)["not_in_base"]
    known = _counts(listed["known"])
    assert (known["probably"], known["similar"], known["none"]) == (1, 1, 2)
    assert _names(similar) == ["15-летний житель Канаша"]
    rows = list(workbook["Результат"].iter_rows(values_only=True))
    column = rows[0].index("В базе Airtable")
    cells = {row[1]: row[column] for row in rows[1:]}
    assert cells["41-летний житель Ленинградской области"] == (
        "вероятно, есть в базе: 41-летний житель <Ленинградской> области"
    )
    assert cells["15-летний житель Канаша"] == (
        "похожие записи в базе: 2 (15-летний житель Канаша; Подросток 15-летний, Канаш (2))"
    )


def test_a_base_of_nameless_records_alone_is_a_base(session_factory: sessionmaker[Session]) -> None:
    _seed(session_factory)
    _unnamed(session_factory, "b", age=15, place="Канаш", articles=["205.5"])
    with _client(session_factory) as client:
        empty = _list(client, months=0, who="unnamed")
        _nameless(session_factory, "15-летний житель Канаша", opened=date(2026, 9, 1))
        loaded = _list(client, months=0, who="unnamed")
        chosen = _list(client, months=0, who="unnamed", known="none")

    # No base at all: nobody is called new. A base of records without a name is a base.
    assert empty["base_loaded"] is False and empty["known"] == []
    assert _row(empty, "Канаша")["known"] is None
    assert loaded["base_loaded"] is True
    assert _row(loaded, "Канаша")["known"]["level"] == "probably"
    assert chosen["items"] == []


def test_the_base_is_asked_with_the_age_of_the_day_it_was_told(
    session_factory: sessionmaker[Session],
) -> None:
    """The age is told on the 5th; a later sentence of the same person, on the 20th, tells
    none. A record «19-летний» of a year ago fits a man of 21 only if the two weeks to
    the later sentence are counted — and they must not be."""
    _seed(session_factory)
    _unnamed(session_factory, "a", age=21, place="Канаш", articles=["205.5"], day=5)
    _unnamed(session_factory, "b", age=None, place="Канаш", articles=["205.5"], day=20)
    _nameless(session_factory, "19-летний житель Канаша", opened=date(2025, 9, 10))

    with _client(session_factory) as client:
        listed = _list(client, months=0, who="unnamed")

    (row,) = listed["items"]  # one person, two sentences
    assert "Канаша" in row["name"]
    assert row["known"] is None and row["not_in_base"] is True
