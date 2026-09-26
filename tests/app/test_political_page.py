"""«Результат»: the politically persecuted off the Rosfinmonitoring list, and its Excel."""

from __future__ import annotations

from collections.abc import Iterator
from contextlib import contextmanager
from datetime import UTC, datetime, timedelta
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session, sessionmaker

from db.orm_models import (
    EntityGroupPoliticsRecord,
    EntityGroupRecord,
    EntityGroupRfMatchRecord,
    RosfinmonitoringEntryRecord,
    RosfinmonitoringSnapshotRecord,
)
from web.app import app
from web.dependencies import get_db


@contextmanager
def _client(session_factory: sessionmaker[Session]) -> Iterator[TestClient]:
    def override_get_db() -> Iterator[Session]:
        with session_factory() as session:
            yield session

    app.dependency_overrides[get_db] = override_get_db
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.pop(get_db, None)


def _seed(session_factory: sessionmaker[Session]) -> None:
    now = datetime.now(UTC)
    with session_factory.begin() as session:
        snapshot = RosfinmonitoringSnapshotRecord(
            snapshot_date=now,
            source_url="https://example.test",
            content_hash="x",
            entry_count=1,
            fetched_at=now,
        )
        session.add(snapshot)
        session.flush()
        entry = RosfinmonitoringEntryRecord(
            snapshot_id=snapshot.id,
            full_name="ИВАНОВ ИВАН ИВАНОВИЧ",
            normalized_name="иванов иван иванович",
            matching_key="ивановиваниванович",
        )
        session.add(entry)
        # Смирнова on the list with her patronymic: who she is, confirmed.
        confirmed = RosfinmonitoringEntryRecord(
            snapshot_id=snapshot.id,
            full_name="СМИРНОВА АННА ПЕТРОВНА",
            normalized_name="смирнова анна петровна",
            matching_key="смирновааннапетровна",
            birth_date=datetime(1990, 2, 1, tzinfo=UTC),
            birth_place="Г. МОСКВА",
        )
        session.add(confirmed)
        for key, name, verdict, days, regions in (
            ("анна смирнова", "Анна Смирнова", "political", 10, [["Москва", 1]]),
            ("иван иванов", "Иван Иванов", "political", 400, []),
            ("александр беда", "Александр Беда", "criminal", 5, []),
        ):
            entity = EntityGroupRecord(
                key=key,
                name=name,
                variants=[[name, 1]],
                mention_count=1,
                article_count=1,
                event_types={},
                regions=regions,
                last_published_at=now - timedelta(days=days),
            )
            session.add(entity)
            session.flush()
            session.add(
                EntityGroupPoliticsRecord(
                    group_id=entity.id,
                    verdict=verdict,
                    method="model",
                    reason=f"так про {name}",
                    quote="цитата",
                )
            )
            if key == "иван иванов":
                session.flush()
                session.add(
                    EntityGroupRfMatchRecord(group_id=entity.id, entry_id=entry.id, level="name")
                )
            if key == "анна смирнова":
                session.flush()
                session.add(
                    EntityGroupRfMatchRecord(
                        group_id=entity.id, entry_id=confirmed.id, level="full"
                    )
                )


def test_the_list_is_the_political_the_period_tells_new_from_old(
    session_factory: sessionmaker[Session],
) -> None:
    _seed(session_factory)

    with _client(session_factory) as client:
        page = client.get("/ui/political").text
        fresh = client.get("/ui/political", params={"months": 3}).text
        certain = client.get("/ui/political", params={"hide_maybe_listed": "true"}).text

    assert "<title>Результат</title>" in page
    assert '<span>Результат</span><span class="nav-count">2</span>' in page
    # The funnel, in one line, down to the result.
    assert "2 политические дела — результат" in page
    assert "Воронка за всё время (публикаций пока нет):" in page
    assert "Найдено: 2." in page and "Беда" not in page
    assert "Смирнова Анна" in page and "Москва" in page and "модель: так про Анна Смирнова" in page
    assert 'Иванов Иван</a> <span class="badge pending">возможно в перечне</span>' in page
    assert "Скрыть возможных в перечне (1)" in page
    # Иванов's latest news is a year old: not a new case.
    assert "Найдено: 1." in fresh and "Иванов" not in fresh
    assert "Найдено: 1." in certain and "Иванов Иван" not in certain
    # On the list with the patronymic: in the result, the list's word beside the name; the
    # box hides only the maybe-namesakes.
    assert 'Смирнова Анна</a> <span class="badge">в перечне РФМ</span>' in page
    assert "СМИРНОВА АННА ПЕТРОВНА, 01.02.1990 г.р., Г. МОСКВА" in page
    assert "Смирнова Анна" in certain


def test_the_excel_has_every_row_of_the_filters(session_factory: sessionmaker[Session]) -> None:
    _seed(session_factory)

    with _client(session_factory) as client:
        response = client.get("/ui/political/export.xlsx", params={"months": 0})

    assert response.status_code == 200
    # All the time: from the earliest latest news (Иванов's, 400 days ago) to today.
    today = datetime.now(UTC).date()
    assert response.headers["content-disposition"] == (
        f'attachment; filename="result_{(today - timedelta(days=400)).isoformat()}_'
        f'{today.isoformat()}.xlsx"'
    )
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:3] == ("№", "Фамилия Имя", "Регион")
    assert rows[0][8] == "Перечень РФМ"
    assert [(row[1], row[2], row[8]) for row in rows[1:]] == [
        (
            "Смирнова Анна",
            "Москва",
            "в перечне: СМИРНОВА АННА ПЕТРОВНА, 01.02.1990 г.р., Г. МОСКВА",
        ),
        ("Иванов Иван", None, "возможно тёзка: ИВАНОВ ИВАН ИВАНОВИЧ"),
    ]


def test_a_period_of_dates_and_the_tick_box(session_factory: sessionmaker[Session]) -> None:
    _seed(session_factory)
    today = datetime.now(UTC).date()
    recent = {"date_from": (today - timedelta(days=30)).isoformat(), "date_to": today.isoformat()}
    old = {
        "date_from": (today - timedelta(days=500)).isoformat(),
        "date_to": (today - timedelta(days=300)).isoformat(),
    }

    with _client(session_factory) as client:
        page = client.get("/ui/political").text
        in_recent = client.get("/ui/political", params=recent).text
        in_old = client.get("/ui/political", params=old).text
        # Dates given, the months are not the choice: a year-old case, though «3 months».
        dates_win = client.get("/ui/political", params={**old, "months": 3}).text
        # The form sends the hidden «false» first, then the ticked box's «true».
        ticked = client.get("/ui/political?hide_maybe_listed=false&hide_maybe_listed=true").text
        excel = client.get("/ui/political/export.xlsx", params=old)

    assert (
        '<input id="box-hide-maybe" type="checkbox" name="hide_maybe_listed" value="true" onchange'
        in page
    )
    assert '<button type="submit" name="months" value="3" class="chip"' in page
    assert '<input type="date" name="date_from" value="">' in page
    # The export sends the form as it is: dates picked without «Показать» count.
    assert '<button type="submit" class="secondary" formaction="/ui/political/export.xlsx">' in page
    assert page.index('<input type="hidden" name="months" value="0">') < page.index(
        'name="months" value="3"'
    )
    # Смирнова's latest news is 10 days old, Иванов's 400.
    assert "Найдено: 1." in in_recent and "Смирнова Анна" in in_recent
    assert "Найдено: 1." in in_old and "Иванов Иван" in in_old
    assert f'name="date_from" value="{old["date_from"]}"' in in_old
    assert "Иванов Иван" in dates_win
    assert "Найдено: 1." in ticked and 'value="true" checked' in ticked
    assert (
        f'filename="result_{old["date_from"]}_{old["date_to"]}.xlsx"'
        in (excel.headers["content-disposition"])
    )
    rows = list(load_workbook(BytesIO(excel.content)).active.iter_rows(values_only=True))  # type: ignore[union-attr]
    assert [row[1] for row in rows[1:]] == ["Иванов Иван"]


def test_the_period_is_kept_for_a_reload_and_the_menu(
    session_factory: sessionmaker[Session],
) -> None:
    _seed(session_factory)
    today = datetime.now(UTC).date()
    old = {
        "date_from": (today - timedelta(days=500)).isoformat(),
        "date_to": (today - timedelta(days=300)).isoformat(),
    }

    with _client(session_factory) as client:
        client.get("/ui/political", params=old)
        # The menu's link, or a reload of it: no filters in the address.
        again = client.get("/ui/political").text
        # «За всё время» chosen is chosen, not the remembered dates.
        client.get("/ui/political", params={"months": 0, "date_from": "", "date_to": ""})
        everything = client.get("/ui/political").text
        # Dates the page's script kept, picked but not shown yet: encoded.
        # In the browser it replaces the server's (same name and path); here, one of two.
        client.cookies.clear()
        client.cookies.set(
            "political_filters",
            f"months%3D0%26date_from%3D{old['date_from']}%26date_to%3D{old['date_to']}"
            "%26hide_maybe_listed%3Dfalse",
        )
        picked = client.get("/ui/political").text

    assert "Найдено: 1." in again and f'name="date_from" value="{old["date_from"]}"' in again
    assert "Найдено: 2." in everything
    assert "Найдено: 1." in picked and "Иванов Иван" in picked
    assert 'id="political-filters"' in picked and "document.cookie" in picked
