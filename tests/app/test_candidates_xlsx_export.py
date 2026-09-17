"""Excel export of the Candidates page (`/ui/candidates/export.xlsx`)."""

from __future__ import annotations

from collections.abc import Iterator
from contextlib import contextmanager
from datetime import UTC, datetime, timedelta
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

import httpx2
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, sessionmaker
from support.research_db_fixtures import ResearchSeeder

from api import app, get_db
from db.orm_models import SourceDocument

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADER = ("№", "Фамилия Имя", "Дата новости", "Категория", "Причины", "Ссылка")
LINK_COLUMN = 6
# The PDF export's «Причины» column: the classifier's reasons joined by «; ».
REASONS = ["Статья содержит признаки политического преследования", "Антивоенная деятельность"]
REASONS_TEXT = "; ".join(REASONS)
# 21:00 UTC is already the next day in Moscow, where these sources publish.
NEWS_TIME = datetime(2026, 8, 31, 21, 0, tzinfo=UTC)
# openpyxl reads a date cell back as a naive datetime.
NEWS_DATE = datetime(2026, 9, 1, tzinfo=UTC).replace(tzinfo=None)


@contextmanager
def _client(session_factory: sessionmaker[Session]) -> Iterator[TestClient]:
    def override_get_db() -> Iterator[Session]:
        with session_factory() as session:
            yield session

    app.dependency_overrides[get_db] = override_get_db
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.pop(get_db, None)


def _export(client: TestClient, snapshot_id: int, **params: Any) -> httpx2.Response:
    """The export with no period filter unless a test sets one."""
    return client.get(
        "/ui/candidates/export.xlsx",
        params={"snapshot_id": snapshot_id, "date_from": "", **params},
    )


def _seed_candidate(
    seed: ResearchSeeder,
    snapshot_id: int,
    name: str,
    *,
    confidence: float = 0.9,
    rf_status: str = "not_matched",
    reasons: list[str] | None = None,
    event_type: str = "arrest",
    published_at: datetime = NEWS_TIME,
) -> int:
    person_id = seed.person(name)
    seed.classification(
        person_id, "political", confidence, reasons=REASONS if reasons is None else reasons
    )
    seed.match(person_id, snapshot_id, rf_status, 0.8)
    _seed_news(seed, person_id, name, f"news-{person_id}", published_at, event_type=event_type)
    return person_id


def _seed_news(
    seed: ResearchSeeder,
    person_id: int,
    name: str,
    external_id: str,
    published_at: datetime,
    *,
    event_type: str = "arrest",
) -> None:
    source_id = seed.source(f"source-{external_id}", f"https://{external_id}.example.test")
    _, run_id = seed.article(
        source_id,
        external_id=external_id,
        title="Новость",
        text=f"Суд арестовал {name}.",
        published_at=published_at,
    )
    seed.event(
        run_id,
        f"Суд арестовал {name}",
        event_type=event_type,
        event_date=published_at,
        links=[(person_id, "subject")],
    )


def _rows(content: bytes) -> list[tuple[object, ...]]:
    sheet = load_workbook(BytesIO(content)).active
    assert sheet is not None
    return list(sheet.iter_rows(values_only=True))


def _news_url(person_id: int) -> str:
    return f"https://example.test/news-{person_id}"


def test_export_returns_downloadable_xlsx_with_candidate_row(
    session_factory: sessionmaker[Session],
) -> None:
    with session_factory() as session:
        seed = ResearchSeeder(session)
        snapshot_id = seed.snapshot()
        person_id = _seed_candidate(seed, snapshot_id, "Иван Иванов")
        session.commit()

    with _client(session_factory) as client:
        response = _export(client, snapshot_id)

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    assert response.headers["content-disposition"] == (
        f'attachment; filename="political-candidates-{snapshot_id}.xlsx"'
    )
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet is not None
    assert sheet.max_column == len(HEADER)
    assert list(sheet.iter_rows(values_only=True)) == [
        HEADER,
        (1, "Иванов Иван", NEWS_DATE, "Арест", REASONS_TEXT, _news_url(person_id)),
    ]
    assert sheet.cell(row=2, column=3).number_format == "DD.MM.YYYY"
    hyperlink = sheet.cell(row=2, column=LINK_COLUMN).hyperlink
    assert hyperlink is not None
    assert hyperlink.target == _news_url(person_id)


def test_names_are_written_surname_first() -> None:
    """Customer request: «сначала фамилию, потом имя», for searching the table."""
    from api import _surname_first

    assert _surname_first("Иван Иванов") == "Иванов Иван"
    assert _surname_first("Владимир Николаевич Казанцев") == "Казанцев Владимир Николаевич"
    # Already surname first: a patronymic closes the name.
    assert _surname_first("Корнилов Алексей Леонидович") == "Корнилов Алексей Леонидович"
    assert _surname_first("Е.А. Аничкина") == "Аничкина Е.А."
    assert _surname_first("Навальный") == "Навальный"


def test_new_cases_and_sentences_come_first_then_the_newest(
    session_factory: sessionmaker[Session],
) -> None:
    """Customer request: new cases and sentences are the priority; the rest follows."""
    with session_factory() as session:
        seed = ResearchSeeder(session)
        snapshot_id = seed.snapshot()
        day = timedelta(days=1)
        _seed_candidate(seed, snapshot_id, "Анна Обыскова", event_type="search")
        _seed_candidate(
            seed, snapshot_id, "Петр Приговоров", event_type="sentence", published_at=NEWS_TIME
        )
        _seed_candidate(
            seed,
            snapshot_id,
            "Олег Делов",
            event_type="case_opened",
            published_at=NEWS_TIME + day,
        )
        _seed_candidate(seed, snapshot_id, "Яков Штрафов", event_type="fine")
        session.commit()

    with _client(session_factory) as client:
        response = _export(client, snapshot_id)

    assert [(row[1], row[3]) for row in _rows(response.content)[1:]] == [
        ("Делов Олег", "Возбуждено дело"),
        ("Приговоров Петр", "Приговор"),
        ("Обыскова Анна", "Обыск"),
        ("Штрафов Яков", "Штраф"),
    ]


def test_export_applies_active_filters(session_factory: sessionmaker[Session]) -> None:
    with session_factory() as session:
        seed = ResearchSeeder(session)
        snapshot_id = seed.snapshot()
        other_snapshot_id = seed.snapshot("other-snapshot-hash")
        strong = _seed_candidate(seed, snapshot_id, "Иван Иванов", confidence=0.95)
        _seed_candidate(seed, snapshot_id, "Слабый Кандидат", confidence=0.75)
        _seed_candidate(seed, snapshot_id, "Есть Вперечне", rf_status="matched")
        _seed_candidate(seed, other_snapshot_id, "Другой Снапшот")
        session.commit()

    with _client(session_factory) as client:
        response = _export(client, snapshot_id, min_confidence=0.9)

    assert _rows(response.content) == [
        HEADER,
        (1, "Иванов Иван", NEWS_DATE, "Арест", REASONS_TEXT, _news_url(strong)),
    ]


def test_old_news_is_left_out_by_the_period(session_factory: sessionmaker[Session]) -> None:
    """Customer finding: news from 2025 filled the end of the table."""
    with session_factory() as session:
        seed = ResearchSeeder(session)
        snapshot_id = seed.snapshot()
        _seed_candidate(seed, snapshot_id, "Иван Свежий")
        _seed_candidate(
            seed, snapshot_id, "Петр Старый", published_at=datetime(2025, 12, 16, tzinfo=UTC)
        )
        session.commit()

    with _client(session_factory) as client:
        filtered = _export(client, snapshot_id, date_from="2026-08-01")
        everything = _export(client, snapshot_id)

    assert [row[1] for row in _rows(filtered.content)[1:]] == ["Свежий Иван"]
    assert [row[1] for row in _rows(everything.content)[1:]] == ["Свежий Иван", "Старый Петр"]


def test_the_default_period_is_the_last_45_days(session_factory: sessionmaker[Session]) -> None:
    today = datetime.now(ZoneInfo("Europe/Moscow"))
    with session_factory() as session:
        seed = ResearchSeeder(session)
        snapshot_id = seed.snapshot()
        _seed_candidate(seed, snapshot_id, "Иван Свежий", published_at=today - timedelta(days=40))
        _seed_candidate(seed, snapshot_id, "Петр Старый", published_at=today - timedelta(days=50))
        session.commit()

    with _client(session_factory) as client:
        response = client.get("/ui/candidates/export.xlsx", params={"snapshot_id": snapshot_id})
        page = client.get("/ui/candidates", params={"snapshot_id": snapshot_id})

    assert [row[1] for row in _rows(response.content)[1:]] == ["Свежий Иван"]
    default_from = (today - timedelta(days=45)).date()
    assert f'name="date_from" value="{default_from.isoformat()}"' in page.text


def test_administrative_cases_are_left_out_unless_asked_for(
    session_factory: sessionmaker[Session],
) -> None:
    """Customer finding: «включает административки, их не надо» (ovdinfolive/42587)."""
    with session_factory() as session:
        seed = ResearchSeeder(session)
        snapshot_id = seed.snapshot()
        _seed_candidate(
            seed,
            snapshot_id,
            "Лена Патяева",
            reasons=["Политическая статья: КоАП РФ ст. 20.2 ч. 8"],
        )
        _seed_candidate(
            seed,
            snapshot_id,
            "Иван Уголовный",
            reasons=["Политическая статья: КоАП РФ ст. 20.3; УК РФ ст. 280.3"],
        )
        # No article at all: the persecution is known from the text, not a КоАП case.
        _seed_candidate(seed, snapshot_id, "Петр Безстатейный")
        session.commit()

    with _client(session_factory) as client:
        criminal = _export(client, snapshot_id)
        everything = _export(client, snapshot_id, include_administrative="1")

    assert [row[1] for row in _rows(criminal.content)[1:]] == [
        "Уголовный Иван",
        "Безстатейный Петр",
    ]
    assert "Патяева Лена" in [row[1] for row in _rows(everything.content)[1:]]


def test_export_contains_all_candidates_not_only_the_page_limit(
    session_factory: sessionmaker[Session],
) -> None:
    with session_factory() as session:
        seed = ResearchSeeder(session)
        snapshot_id = seed.snapshot()
        names = ["Первый Кандидат", "Второй Кандидат", "Третий Кандидат"]
        for name in names:
            _seed_candidate(seed, snapshot_id, name)
        session.commit()

    params: dict[str, str | int] = {"snapshot_id": snapshot_id, "limit": 2, "date_from": ""}
    with _client(session_factory) as client:
        page = client.get("/ui/candidates", params=params)
        response = client.get("/ui/candidates/export.xlsx", params=params)

    assert page.text.count("Кандидат</a>") == 2
    assert len(_rows(response.content)) == 1 + len(names)


def test_export_without_candidates_has_only_header(
    session_factory: sessionmaker[Session],
) -> None:
    with session_factory() as session:
        snapshot_id = ResearchSeeder(session).snapshot()
        session.commit()

    with _client(session_factory) as client:
        response = _export(client, snapshot_id)

    assert response.status_code == 200
    assert _rows(response.content) == [HEADER]


def test_export_of_unknown_snapshot_is_not_found(
    session_factory: sessionmaker[Session],
) -> None:
    with _client(session_factory) as client:
        response = _export(client, 999)

    assert response.status_code == 404


def test_candidates_page_links_export_with_active_filters(
    session_factory: sessionmaker[Session],
) -> None:
    with session_factory() as session:
        snapshot_id = ResearchSeeder(session).snapshot()
        session.commit()

    with _client(session_factory) as client:
        page = client.get(
            "/ui/candidates",
            params={
                "snapshot_id": snapshot_id,
                "min_confidence": 0.8,
                "date_from": "2026-08-01",
                "include_administrative": "1",
            },
        )

    assert (
        f'href="/ui/candidates/export.xlsx?snapshot_id={snapshot_id}&min_confidence=0.8'
        '&date_from=2026-08-01&include_administrative=1"' in page.text
    )
    assert ">Export to Excel</a>" in page.text


def test_export_writes_formula_like_name_as_text(
    session_factory: sessionmaker[Session],
) -> None:
    name = '=HYPERLINK("https://evil.test","Иван")'
    with session_factory() as session:
        seed = ResearchSeeder(session)
        snapshot_id = seed.snapshot()
        _seed_candidate(seed, snapshot_id, name)
        session.commit()

    with _client(session_factory) as client:
        response = _export(client, snapshot_id)

    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet is not None
    name_cell = sheet.cell(row=2, column=2)
    assert name_cell.value is not None and "HYPERLINK" in str(name_cell.value)
    assert name_cell.data_type == "s"


def test_link_points_to_the_article_of_the_latest_event(
    session_factory: sessionmaker[Session],
) -> None:
    with session_factory() as session:
        seed = ResearchSeeder(session)
        snapshot_id = seed.snapshot()
        person_id = _seed_candidate(seed, snapshot_id, "Иван Иванов")
        _seed_news(seed, person_id, "Иван Иванов", "later-news", datetime(2026, 9, 10, tzinfo=UTC))
        _seed_news(seed, person_id, "Иван Иванов", "older-news", datetime(2026, 8, 1, tzinfo=UTC))
        session.commit()

    with _client(session_factory) as client:
        response = _export(client, snapshot_id)

    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet is not None
    assert sheet.cell(row=2, column=LINK_COLUMN).value == "https://example.test/later-news"
    assert sheet.cell(row=2, column=3).value == datetime(2026, 9, 10, tzinfo=UTC).replace(
        tzinfo=None
    )
    hyperlink = sheet.cell(row=2, column=LINK_COLUMN).hyperlink
    assert hyperlink is not None
    assert hyperlink.target == "https://example.test/later-news"


def test_link_without_events_points_to_the_article_that_mentions_the_person(
    session_factory: sessionmaker[Session],
) -> None:
    with session_factory() as session:
        seed = ResearchSeeder(session)
        snapshot_id = seed.snapshot()
        person_id = seed.person("Иван Иванов")
        seed.classification(person_id, "political", 0.9)
        seed.match(person_id, snapshot_id, "not_matched", 0.8)
        source_id = seed.source("mention-source", "https://example.test")
        _, run_id = seed.article(
            source_id,
            external_id="mention-news",
            title="Пикет",
            text="Иван Иванов вышел.",
            published_at=NEWS_TIME,
        )
        seed.mention(run_id, "Иван Иванов", person_id=person_id)
        session.commit()

    with _client(session_factory) as client:
        response = _export(client, snapshot_id)

    assert _rows(response.content) == [
        HEADER,
        (1, "Иванов Иван", NEWS_DATE, None, None, "https://example.test/mention-news"),
    ]


def test_only_a_web_link_is_clickable(session_factory: sessionmaker[Session]) -> None:
    with session_factory() as session:
        seed = ResearchSeeder(session)
        snapshot_id = seed.snapshot()
        person_id = _seed_candidate(seed, snapshot_id, "Иван Иванов")
        document = session.scalars(
            select(SourceDocument).where(SourceDocument.external_id == f"news-{person_id}")
        ).one()
        document.canonical_url = "javascript:alert(1)"
        session.commit()

    with _client(session_factory) as client:
        response = _export(client, snapshot_id)

    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet is not None
    assert sheet.cell(row=2, column=LINK_COLUMN).value == "javascript:alert(1)"
    assert sheet.cell(row=2, column=LINK_COLUMN).hyperlink is None
